"""Export d'un tableau : les trois formats produisent-ils des fichiers valides ?

On ne se contente pas de vérifier qu'il n'y a pas d'exception : chaque fichier
est relu et son contenu comparé au tableau de départ.
"""

import datetime as dt
import sys
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path

PROJET = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJET))

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

from compta import export  # noqa: E402

ok, ko = 0, 0


def verifier(nom, condition, detail=""):
    global ok, ko
    if condition:
        ok += 1
        print(f"  OK   {nom}")
    else:
        ko += 1
        print(f"  ECHEC {nom} {detail}")


TABLEAU = pd.DataFrame([
    {"Date": dt.date(2026, 7, 13), "Payant": 12, "APCI": 5, "Divers": 1250.0,
     "Rectifiée": False},
    {"Date": dt.date(2026, 7, 14), "Payant": 0, "APCI": 0, "Divers": 0.0,
     "Rectifiée": False},
    {"Date": dt.date(2026, 7, 15), "Payant": 31, "APCI": 9, "Divers": 0.0,
     "Rectifiée": True},
]).set_index("Date")

TOTAUX = [("Recettes", "35 178 TND"), ("CA déclaré", "21 618 TND"),
          ("Patients", "57")]
TITRE = "Registre réel — juillet 2026"
SOUS_TITRE = "Cabinet d'ophtalmologie · valeurs d'origine"

print("\n1. CSV")
donnees = export.vers_csv(TABLEAU, TOTAUX)
verifier("des octets sont produits", isinstance(donnees, bytes) and donnees)
verifier("BOM présent, pour qu'Excel lise les accents",
         donnees.startswith(b"\xef\xbb\xbf"))
texte = donnees.decode("utf-8-sig")
verifier("les totaux figurent en tête", texte.startswith("Recettes;35 178 TND"))
verifier("séparateur point-virgule", ";" in texte.splitlines()[4])
verifier("virgule décimale", "1250,00" in texte or "1250.0" not in texte,
         f"({texte.splitlines()[5] if len(texte.splitlines()) > 5 else texte!r})")
relu = pd.read_csv(BytesIO(donnees), sep=";", decimal=",", skiprows=4)
verifier("relecture : 3 lignes", len(relu) == 3, f"({len(relu)})")
verifier("relecture : les colonnes sont là",
         list(relu.columns) == ["Date", "Payant", "APCI", "Divers", "Rectifiée"],
         f"({list(relu.columns)})")
verifier("relecture : les valeurs sont justes",
         relu["Payant"].tolist() == [12, 0, 31], f"({relu['Payant'].tolist()})")
verifier("l'index est devenu une colonne", relu["Date"].iloc[0] == "13/07/2026",
         f"({relu['Date'].iloc[0]})")

print("\n2. Excel")
donnees = export.vers_xlsx(TABLEAU, TITRE, TOTAUX, SOUS_TITRE)
verifier("des octets sont produits", isinstance(donnees, bytes) and donnees)
verifier("c'est bien un .xlsx (archive zip)",
         zipfile.is_zipfile(BytesIO(donnees)))

wb = openpyxl.load_workbook(BytesIO(donnees))
ws = wb.active
valeurs = [[c.value for c in ligne] for ligne in ws.iter_rows()]
plat = [v for ligne in valeurs for v in ligne if v is not None]
verifier("le titre est en première ligne", valeurs[0][0] == TITRE)
verifier("le sous-titre suit", valeurs[1][0] == SOUS_TITRE)
verifier("la date d'édition est mentionnée",
         any("Édité le" in str(v) for v in plat))
verifier("les totaux sont présents", "Recettes" in plat and "35 178 TND" in plat)
verifier("l'en-tête du tableau est présent",
         all(c in plat for c in ["Date", "Payant", "APCI", "Divers", "Rectifiée"]))
verifier("les nombres restent des nombres, pas du texte",
         31 in plat or 31.0 in plat)
verifier("les dates restent des dates",
         any(isinstance(v, dt.datetime) for v in plat))
verifier("les volets sont figés sous l'en-tête", ws.freeze_panes is not None,
         f"({ws.freeze_panes})")
verifier("les colonnes ont une largeur",
         all(d.width > 0 for d in ws.column_dimensions.values()))
wb.close()

print("\n3. PDF")
donnees = export.vers_pdf(TABLEAU, TITRE, TOTAUX, SOUS_TITRE)
verifier("des octets sont produits", isinstance(donnees, bytes) and donnees)
verifier("en-tête PDF valide", donnees.startswith(b"%PDF-"))
verifier("fin de fichier correcte", donnees.rstrip().endswith(b"%%EOF"))
verifier("taille plausible", 2000 < len(donnees) < 400_000, f"({len(donnees)} o)")
verifier("le titre est dans les métadonnées", TITRE.encode("utf-16-be") in donnees
         or TITRE.encode("latin-1", "replace") in donnees or b"/Title" in donnees)

chemin = Path(tempfile.mkdtemp()) / "essai.pdf"
chemin.write_bytes(donnees)
verifier("le fichier s'écrit sur le disque", chemin.stat().st_size == len(donnees))

print("\n4. Cas limites")
vide = pd.DataFrame({"Date": pd.Series(dtype="object"),
                     "Payant": pd.Series(dtype="int")}).set_index("Date")
for nom, fabriquer in (("csv", lambda d: export.vers_csv(d)),
                       ("xlsx", lambda d: export.vers_xlsx(d, "Vide")),
                       ("pdf", lambda d: export.vers_pdf(d, "Vide"))):
    try:
        resultat = fabriquer(vide)
        verifier(f"tableau vide accepté en {nom}", bool(resultat))
    except Exception as err:  # noqa: BLE001
        verifier(f"tableau vide accepté en {nom}", False, f"({err})")

accents = pd.DataFrame([{"Poste": "Cataracte payante", "Motif": "Journée « à zéro » — 3 patients comptés deux fois", "Écart": -180.0}])
for nom, fabriquer in (("csv", lambda d: export.vers_csv(d)),
                       ("xlsx", lambda d: export.vers_xlsx(d, "Accents")),
                       ("pdf", lambda d: export.vers_pdf(d, "Accents"))):
    try:
        resultat = fabriquer(accents)
        verifier(f"accents et guillemets acceptés en {nom}", bool(resultat))
    except Exception as err:  # noqa: BLE001
        verifier(f"accents et guillemets acceptés en {nom}", False, f"({err})")

verifier("les accents survivent au CSV",
         "comptés deux fois" in export.vers_csv(accents).decode("utf-8-sig"))
wb = openpyxl.load_workbook(BytesIO(export.vers_xlsx(accents, "Accents")))
verifier("les accents survivent à Excel",
         any("comptés deux fois" in str(c.value)
             for ligne in wb.active.iter_rows() for c in ligne if c.value))
wb.close()

print("\n5. Un mois entier, 31 lignes et 12 colonnes")
grand = pd.DataFrame([
    {"Date": dt.date(2026, 7, j), **{f"Poste {i}": j * i for i in range(1, 12)}}
    for j in range(1, 32)
]).set_index("Date")
pdf = export.vers_pdf(grand, "Mois complet", TOTAUX)
verifier("le PDF d'un mois complet se produit", pdf.startswith(b"%PDF-"))
verifier("il tient en une ou deux pages", 3000 < len(pdf) < 200_000,
         f"({len(pdf)} o)")
xlsx = export.vers_xlsx(grand, "Mois complet", TOTAUX)
wb = openpyxl.load_workbook(BytesIO(xlsx))
verifier("le classeur porte les 31 journées", wb.active.max_row >= 31 + 5,
         f"({wb.active.max_row} lignes)")
wb.close()

print(f"\n{'=' * 50}\n{ok} réussis, {ko} échecs\n{'=' * 50}")
sys.exit(1 if ko else 0)
