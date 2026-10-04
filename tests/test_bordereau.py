"""Le bordereau mensuel : structure, formules, tarifs, archivage, historique.

Tout se joue sur une COPIE de la base et dans un dossier temporaire — le
registre du cabinet et le classeur ne sont jamais touchés.

Ce qui compte ici et qu'aucun autre test ne couvre : le bordereau doit sortir
avec des **formules vivantes**, au bon numéro de ligne, et chaque journée doit
porter le tarif de son époque. Un bordereau qui figerait des nombres se lirait
pareil et ne vaudrait plus rien.
"""

import datetime as dt
import shutil
import sys
import tempfile
from io import BytesIO
from pathlib import Path

PROJET = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJET))

# le cabinet fictif de tests/atelier.py : ni le classeur réel, ni data/compta.db
import atelier  # noqa: E402

CLASSEUR, BASE = atelier.installer()

import openpyxl  # noqa: E402

from compta import bordereau, store  # noqa: E402
from compta.model import Journee, TARIFS_DEFAUT  # noqa: E402

tmp = Path(tempfile.mkdtemp())
base = tmp / "test.db"
shutil.copy2(BASE, base)
store.DB_PATH = base

conn = store.connect()
store.init_db(conn)

ok, ko = 0, 0


def verifier(nom, condition, detail=""):
    global ok, ko
    if condition:
        ok += 1
        print(f"  OK   {nom}")
    else:
        ko += 1
        print(f"  ECHEC {nom} {detail}")


def feuille(contenu: bytes, annee: int, mois: int):
    wb = openpyxl.load_workbook(BytesIO(contenu))
    return wb, wb[bordereau.nom_feuille(annee, mois)]


print("\n1. Le mois en journées")
aout = store.mois_en_journees(conn, 2026, 8)
verifier("août 2026 fait 31 journées", len(aout) == 31, f"({len(aout)})")
verifier("elles vont du 1er au 31",
         aout[0].date == dt.date(2026, 8, 1) and aout[-1].date == dt.date(2026, 8, 31))
verifier("elles sont dans l'ordre",
         all(a.date < b.date for a, b in zip(aout, aout[1:])))
fevrier = store.mois_en_journees(conn, 2024, 2)
verifier("février bissextile fait 29 journées", len(fevrier) == 29,
         f"({len(fevrier)})")
# le classeur s'arrête au 23/01/2027 : un mois au-delà se construit quand même
futur = store.mois_en_journees(conn, 2027, 6)
verifier("un mois absent du registre fait quand même ses 30 lignes",
         len(futur) == 30, f"({len(futur)})")
verifier("ses journées sont à zéro", all(j.vide for j in futur))

print("\n2. Structure du classeur produit")
contenu = bordereau.construire(aout, 2026, 8, registre="corrigé")
verifier("des octets sont produits", isinstance(contenu, bytes) and contenu)
wb, ws = feuille(contenu, 2026, 8)
verifier("la feuille du mois porte son nom", ws.title == "août 2026", f"({ws.title})")
verifier("une feuille Récapitulatif l'accompagne", "Récapitulatif" in wb.sheetnames,
         f"({wb.sheetnames})")
verifier("les colonnes vont de A à AD", ws.max_column == 30, f"({ws.max_column})")
verifier("31 journées + 3 lignes d'en-tête + 1 total = 35 lignes",
         ws.max_row == 35, f"({ws.max_row})")
verifier("les journées commencent en ligne 4",
         bordereau.PREMIERE_LIGNE == 4 and isinstance(ws["B4"].value, dt.datetime))
verifier("l'en-tête est celui du classeur",
         ws["B1"].value == "DATE" and ws["C1"].value == "C° CNAM"
         and ws["U1"].value == "RECETTES" and ws["AD1"].value == "TVA DED")
verifier("la première journée est le samedi 1er août",
         ws["A4"].value == "SAMEDI"
         and ws["B4"].value.date() == dt.date(2026, 8, 1))
verifier("les volets sont figés sous l'en-tête", ws.freeze_panes == "C4",
         f"({ws.freeze_panes})")
verifier("l'en-tête se répète à l'impression",
         (ws.print_title_rows or "").replace("$", "") == "1:3",
         f"({ws.print_title_rows})")

print("\n3. Les formules restent des formules")
attendu = {
    "D": "=C4*72", "F": "=E4*20", "G": "=E4*21", "I": "=H4*70", "K": "=J4*70",
    "M": "=L4*70", "O": "=N4*700", "Q": "=P4*750", "W": "=E4*11",
    "U": "=D4+F4+I4+K4+M4+O4+Q4+S4+T4+R4+W4",
    "X": "=SUM(F4,I4,M4,T4,W4)-V4", "Y": "=K4", "Z": "=SUM(D4,G4,O4)",
    "AA": "=SUM(F4,I4,K4,M4,O4/2)", "AB": "=C4+E4+H4+J4+L4",
    "AC": "=D4+F4+G4+I4+K4+M4+O4+R4", "AD": "=AC4-(AC4/1.06)",
}
for colonne, formule in attendu.items():
    verifier(f"{colonne}4 est la formule du classeur",
             ws[f"{colonne}4"].value == formule, f"({ws[f'{colonne}4'].value!r})")

verifier("les formules suivent le numéro de ligne",
         ws["U34"].value == "=D34+F34+I34+K34+M34+O34+Q34+S34+T34+R34+W34",
         f"({ws['U34'].value!r})")
verifier("aucune colonne calculée n'est figée en nombre",
         all(isinstance(ws[f"{c}{r}"].value, str)
             and ws[f"{c}{r}"].value.startswith("=")
             for c in attendu for r in range(4, 35)))

compteurs = {"C": "apci_n", "E": "tp_n", "H": "remb_n", "J": "star_n",
             "L": "pay_n", "N": "cat_cnam_n", "P": "cat_pay_n"}
verifier("les compteurs sont écrits en valeur, pas en formule",
         all(isinstance(ws[f"{c}4"].value, (int, float)) for c in compteurs))
verifier("ce sont bien les compteurs du registre",
         all(ws[f"{col}4"].value == getattr(aout[0], cle)
             for col, cle in compteurs.items()),
         f"({[(c, ws[f'{c}4'].value) for c in compteurs]})")
verifier("les montants saisis aussi — R, S, T, V",
         all(isinstance(ws[f"{c}4"].value, (int, float)) for c in "RSTV"))

print("\n4. Chaque journée porte le tarif de son époque")
# 2018 : le payant était à 35 ou 60 TND, jamais à celui d'aujourd'hui
mars_2018 = store.mois_en_journees(conn, 2018, 3)
tarifs_2018 = bordereau.tarifs_du_mois(mars_2018)
verifier("2018 ne s'écrit pas au tarif de 2026",
         TARIFS_DEFAUT["t_pay"] not in tarifs_2018["t_pay"],
         f"(payant 2018 : {tarifs_2018['t_pay']}, aujourd'hui "
         f"{TARIFS_DEFAUT['t_pay']})")
vieux = bordereau.construire(mars_2018, 2018, 3)
_, ws_2018 = feuille(vieux, 2018, 3)
tarif_ligne4 = mars_2018[0].tarifs["t_pay"]
verifier("la formule du bordereau 2018 porte le tarif de 2018",
         ws_2018["M4"].value == f"=L4*{int(tarif_ligne4)}",
         f"({ws_2018['M4'].value!r}, attendu =L4*{int(tarif_ligne4)})")
verifier("le récapitulatif 2018 liste ses tarifs, pas ceux d'aujourd'hui",
         any(str(int(tarif_ligne4)) in str(c.value)
             for ligne in openpyxl.load_workbook(BytesIO(vieux))["Récapitulatif"]
             .iter_rows() for c in ligne if c.value))

# un tarif qui change en cours de mois doit se voir, pas se moyenner
melange = [
    Journee(date=dt.date(2026, 1, 1), pay_n=1, tarifs=dict(TARIFS_DEFAUT) | {"t_pay": 60.0}),
    Journee(date=dt.date(2026, 1, 2), pay_n=1, tarifs=dict(TARIFS_DEFAUT) | {"t_pay": 70.0}),
]
verifier("deux tarifs dans le mois sont tous deux relevés",
         bordereau.tarifs_du_mois(melange)["t_pay"] == [60.0, 70.0],
         f"({bordereau.tarifs_du_mois(melange)['t_pay']})")
_, ws_melange = feuille(bordereau.construire(melange, 2026, 1), 2026, 1)
verifier("chaque ligne garde le sien",
         ws_melange["M4"].value == "=L4*60" and ws_melange["M5"].value == "=L5*70",
         f"({ws_melange['M4'].value!r}, {ws_melange['M5'].value!r})")

print("\n5. La ligne de total additionne exactement les journées")
total = 35
verifier("elle suit la dernière journée", ws[f"A{total}"].value == "TOTAL",
         f"({ws[f'A{total}'].value!r})")
verifier("elle rappelle le mois", ws[f"B{total}"].value == "août 2026",
         f"({ws[f'B{total}'].value!r})")
for colonne in ("C", "L", "U", "AB", "AC", "AD"):
    verifier(f"{colonne}{total} additionne les lignes 4 à 34",
             ws[f"{colonne}{total}"].value == f"=SUM({colonne}4:{colonne}34)",
             f"({ws[f'{colonne}{total}'].value!r})")
verifier("les 28 colonnes chiffrées ont leur total",
         all(ws[f"{c}{total}"].value == f"=SUM({c}4:{c}34)"
             for c in bordereau.COLONNES))
verifier("le total ne déborde pas sur l'en-tête",
         not any(ws[f"{c}{total}"].value.endswith(":C3)")
                 for c in bordereau.COLONNES))

print("\n6. Le récapitulatif renvoie au mois, il ne le recopie pas")
rec = wb["Récapitulatif"]
valeurs = {ligne[0].value: ligne[1].value for ligne in rec.iter_rows(max_col=2)
           if ligne[0].value}
verifier("il annonce le mois", "août 2026" in str(rec["A1"].value),
         f"({rec['A1'].value!r})")
verifier("il dit de quel registre il vient",
         valeurs.get("Registre") == "registre corrigé",
         f"({valeurs.get('Registre')!r})")
verifier("il date l'édition", "Édité le" in valeurs)
verifier("les recettes renvoient à la ligne de total du mois",
         valeurs.get("Recettes") == "='août 2026'!U35",
         f"({valeurs.get('Recettes')!r})")
verifier("le CA à déclarer renvoie à la colonne AC",
         valeurs.get("CA à déclarer") == "='août 2026'!AC35",
         f"({valeurs.get('CA à déclarer')!r})")
verifier("« à verser à la banque » y figure et renvoie à la colonne AA",
         valeurs.get("À verser à la banque") == "='août 2026'!AA35",
         f"({valeurs.get('À verser à la banque')!r})")
verifier("son en-tête est en place sur la feuille du mois",
         (ws["AA1"].value, ws["AA2"].value, ws["AA3"].value)
         == ("a verser", "a", "la banque"),
         f"({ws['AA1'].value!r}, {ws['AA2'].value!r}, {ws['AA3'].value!r})")
verifier("aucun total n'est recopié en dur",
         all(isinstance(valeurs.get(libelle), str)
             and str(valeurs.get(libelle)).startswith("=")
             for _, libelle, _ in bordereau.POSTES_TOTAUX))

print("\n6 bis. Rien ne se perd entre le fichier et ce qu'on en montre")
# Le garde-fou : l'aperçu affiché dans l'app doit couvrir exactement les
# colonnes que le fichier emporte. Sans ce test, ajouter une colonne au
# bordereau sans l'ajouter à l'aperçu passerait inaperçu — c'est exactement
# comme ça que « à verser à la banque » avait disparu de l'écran.
lettres_apercu = [lettre for lettre, _, _ in bordereau.APERCU]
verifier("l'aperçu couvre les 28 colonnes chiffrées, dans le même ordre",
         lettres_apercu == bordereau.COLONNES,
         f"(manquantes : {sorted(set(bordereau.COLONNES) - set(lettres_apercu))}, "
         f"en trop : {sorted(set(lettres_apercu) - set(bordereau.COLONNES))})")
verifier("chaque colonne de l'aperçu porte un libellé distinct",
         len({libelle for _, libelle, _ in bordereau.APERCU}) == len(bordereau.APERCU))
verifier("chaque colonne de l'aperçu sait où lire sa valeur",
         all(hasattr(aout[0], champ) for _, _, champ in bordereau.APERCU),
         f"({[c for _, _, c in bordereau.APERCU if not hasattr(aout[0], c)]})")
verifier("les compteurs de l'aperçu sont bien des compteurs",
         bordereau.APERCU_COMPTEURS == {"APCI", "TP", "Remb", "Assurance",
                                        "Payant", "Cat. CNAM", "Cat. payante",
                                        "Patients"},
         f"({sorted(bordereau.APERCU_COMPTEURS)})")
# et le récapitulatif ne doit oublier aucun poste de recette
postes = {colonne for _, _, colonne in bordereau.POSTES_TOTAUX}
verifier("le récapitulatif reprend les recettes, le CA, la caisse, le dû CNAM, "
         "la banque et la TVA",
         {"U", "AC", "X", "Z", "AA", "AD"} <= postes, f"({sorted(postes)})")
verifier("il reprend aussi le laser, le divers et les dépenses",
         {"R", "S", "T", "V"} <= postes, f"({sorted(postes)})")

print("\n7. Les totaux annoncés sont ceux des journées")
t = bordereau.totaux(aout)
verifier("les recettes sont la somme des journées",
         abs(t["recettes"] - sum(j.recettes for j in aout)) < 0.005)
verifier("le CA à déclarer aussi",
         abs(t["ca_declare"] - sum(j.ca_declare for j in aout)) < 0.005)
verifier("les patients sont un entier de journées",
         t["patients"] == sum(j.patients for j in aout))
verifier("le non déclaré est l'écart des deux",
         abs(t["non_declare"] - (t["recettes"] - t["ca_declare"])) < 0.005)
verifier("les jours actifs ne comptent que les journées qui rapportent",
         t["jours_actifs"] == sum(1 for j in aout if j.recettes > 0))
verifier("le mois entier est compté", t["jours"] == 31)
verifier("le à-verser-à-la-banque est de la partie",
         abs(t["banque"] - sum(j.banque for j in aout)) < 0.005)
# les agrégats mensuels le portent aussi, sinon l'écart d'un mois à l'autre ne
# s'affiche pas sur la page
agregats = {r["mois"]: r for r in store.agregat_mensuel(conn)}
verifier("les agrégats mensuels portent la banque",
         "banque" in agregats["2026-08"], f"({sorted(agregats['2026-08'])})")
verifier("et ils y retrouvent le même total",
         abs(float(agregats["2026-08"]["banque"]) - t["banque"]) < 0.5,
         f"({agregats['2026-08']['banque']} vs {t['banque']})")

print("\n8. Enregistrer : l'édition précédente est archivée, pas écrasée")
dossier = tmp / "bordereaux"
chemin1, archive1 = bordereau.enregistrer(contenu, 2026, 8, dossier=dossier)
verifier("le fichier est écrit", chemin1.exists() and chemin1.stat().st_size > 0)
verifier("il est rangé par année", chemin1.parent.name == "2026",
         f"({chemin1.parent})")
verifier("son nom dit le mois", chemin1.name == "bordereau-2026-08-aout.xlsx",
         f"({chemin1.name})")
verifier("la première édition n'archive rien", archive1 is None)

autre = bordereau.construire(aout[:20], 2026, 8)
chemin2, archive2 = bordereau.enregistrer(autre, 2026, 8, dossier=dossier)
verifier("la seconde édition reprend le même chemin", chemin2 == chemin1)
verifier("elle archive la première", archive2 is not None and archive2.exists())
verifier("l'archive porte l'ancien contenu",
         archive2.read_bytes() == contenu, "(contenu différent)")
verifier("le chemin courant porte la nouvelle édition",
         chemin2.read_bytes() == autre)
verifier("l'archive est horodatée", "_20" in archive2.name, f"({archive2.name})")

troisieme, archive3 = bordereau.enregistrer(contenu, 2026, 8, dossier=dossier)
verifier("une troisième édition archive la deuxième, sans écraser la première",
         archive3 != archive2 and archive2.exists() and archive3.exists())

print("\n9. L'historique des éditions")
avant = len(store.liste_bordereaux(conn))
identifiant = store.enregistrer_bordereau(
    conn, 2026, 8, "corrigé", bordereau.empreinte(contenu), t,
    fichier=str(chemin1))
verifier("l'édition est consignée", identifiant > 0)
editions = store.bordereaux_du_mois(conn, 2026, 8)
verifier("elle se relit sur son mois", len(editions) == 1, f"({len(editions)})")
verifier("ses totaux sont rendus en dictionnaire",
         isinstance(editions[0]["totaux"], dict)
         and abs(editions[0]["totaux"]["recettes"] - t["recettes"]) < 0.005)
verifier("son empreinte est conservée",
         editions[0]["empreinte"] == bordereau.empreinte(contenu))
verifier("le chemin du fichier est noté", editions[0]["fichier"] == str(chemin1))

# une réédition s'ajoute, elle ne remplace pas
store.enregistrer_bordereau(conn, 2026, 8, "réel",
                            bordereau.empreinte(autre),
                            bordereau.totaux(aout[:20]))
editions = store.bordereaux_du_mois(conn, 2026, 8)
verifier("rééditer ajoute une ligne", len(editions) == 2, f"({len(editions)})")
verifier("la plus récente vient en tête", editions[0]["registre"] == "réel",
         f"({editions[0]['registre']})")
verifier("l'édition précédente est intacte",
         abs(editions[1]["totaux"]["recettes"] - t["recettes"]) < 0.005)
verifier("l'écart entre deux éditions se lit",
         editions[0]["totaux"]["recettes"] < editions[1]["totaux"]["recettes"])

store.enregistrer_bordereau(conn, 2026, 7, "corrigé", "abc123",
                            bordereau.totaux(store.mois_en_journees(conn, 2026, 7)))
derniers = store.derniers_bordereaux_par_mois(conn)
# la base de travail peut déjà porter de vrais bordereaux — l'app sert : on
# vérifie ce que ce test a ajouté, pas l'état absolu de la table
verifier("chaque mois n'apparaît qu'une fois dans la vue d'évolution",
         {"2026-07", "2026-08"} <= set(derniers)
         and len(derniers) == len(set(derniers)),
         f"({sorted(derniers)})")
verifier("les trois éditions de ce test se replient sur deux mois",
         len([m for m in derniers if m in {"2026-07", "2026-08"}]) == 2,
         f"({sorted(derniers)})")
verifier("c'est la dernière édition du mois qui y figure",
         derniers["2026-08"]["registre"] == "réel")
verifier("la liste complète garde toutes les éditions",
         len(store.liste_bordereaux(conn)) == avant + 3,
         f"({len(store.liste_bordereaux(conn))})")
verifier("elle se limite à la demande",
         len(store.liste_bordereaux(conn, limite=2)) == 2)

print("\n10. Cas limites")
try:
    bordereau.construire([], 2026, 8)
    verifier("un mois sans journée est refusé", False, "(aucune erreur)")
except ValueError:
    verifier("un mois sans journée est refusé", True)

vide = store.mois_en_journees(conn, 2027, 6)
contenu_vide = bordereau.construire(vide, 2027, 6)
_, ws_vide = feuille(contenu_vide, 2027, 6)
verifier("un mois à zéro produit quand même ses formules",
         ws_vide["U4"].value.startswith("=") and ws_vide["C4"].value == 0)
verifier("et sa ligne de total", ws_vide["A34"].value == "TOTAL",
         f"({ws_vide['A34'].value!r})")

verifier("le nom de fichier de février est lisible",
         bordereau.nom_fichier(2026, 2) == "bordereau-2026-02-fevrier.xlsx",
         f"({bordereau.nom_fichier(2026, 2)})")
verifier("celui d'août n'a pas d'accent dans le nom",
         bordereau.nom_fichier(2026, 8) == "bordereau-2026-08-aout.xlsx",
         f"({bordereau.nom_fichier(2026, 8)})")
verifier("décembre et janvier ne se mélangent pas",
         bordereau.nom_fichier(2026, 12) < bordereau.nom_fichier(2027, 1))

decembre = store.mois_en_journees(conn, 2026, 12)
verifier("décembre fait 31 journées", len(decembre) == 31, f"({len(decembre)})")
verifier("le mois de 30 jours en fait 30",
         len(store.mois_en_journees(conn, 2026, 11)) == 30)

print("\n11. Le dossier du comptable — noms de dossiers, lecture, rangement")
from compta import dossier_comptable as dc  # noqa: E402
from openpyxl import Workbook  # noqa: E402

# on reconnaît le mois malgré les noms réels du dossier, fautes comprises
for nom, attendu in [("aout 2026", 8), ("juuillet 2014", 7), ("ocobre 2024", 10),
                     ("juin2016", 6), ("fevrier  2016", 2), ("aout", 8),
                     ("MARS 2012", 3), ("septembre", 9), ("decembre 2013", 12)]:
    verifier(f"« {nom} » désigne le mois {attendu}",
             dc.mois_du_nom(nom) == attendu, f"({dc.mois_du_nom(nom)})")
for nom in ("factures avril 2019", "STEG", "2015_03_16", "retenue",
            "facture nov 2018", "declaration janvier 2012", "TELECOM"):
    verifier(f"« {nom} » n'est pas un dossier de mois",
             dc.mois_du_nom(nom) is None, f"({dc.mois_du_nom(nom)})")

# une arborescence fabriquée, à l'image de la vraie
faux = tmp / "comptabilite"
(faux / "2026" / "aout 2026").mkdir(parents=True)
(faux / "2026" / "juuillet 2026").mkdir(parents=True)
(faux / "2026" / "factures mai 2026").mkdir(parents=True)


def fabriquer_transmis(chemin, annee, mois, jours, ca_jour, patients_jour):
    """Un classeur mensuel comme ceux du comptable : en-têtes sur trois lignes,
    valeurs en dur (c'est ce qu'Excel met en cache), colonnes décalées."""
    wb = Workbook()
    ws = wb.active
    ws["B1"] = "DATE"
    ws["H1"] = "CA A "
    ws["H2"] = "DECLARER"
    ws["I1"] = "TVA DED"
    ws["G2"] = "patients à "
    ws["G3"] = "DECLARER"
    for n in range(jours):
        r = 4 + n
        ws.cell(row=r, column=2, value=dt.datetime(annee, mois, n + 1))
        ws.cell(row=r, column=7, value=patients_jour)
        ws.cell(row=r, column=8, value=ca_jour)
        ws.cell(row=r, column=9, value=round(ca_jour * 0.0566, 2))
    wb.save(chemin)
    wb.close()


transmis_aout = faux / "2026" / "aout 2026" / "aout 2026.xlsx"
fabriquer_transmis(transmis_aout, 2026, 8, 31, 500.0, 5)
fabriquer_transmis(faux / "2026" / "juuillet 2026" / "juillet 2026.xlsx",
                   2026, 7, 31, 400.0, 4)
# un fichier qui n'est pas une feuille mensuelle : il doit être écarté
wb_bruit = Workbook()
wb_bruit.active["A1"] = "facture STEG"
wb_bruit.save(faux / "2026" / "factures mai 2026" / "steg.xlsx")
wb_bruit.close()

verifier("le dossier d'un mois se retrouve par son nom exact",
         dc.dossier_du_mois(2026, 8, faux) == faux / "2026" / "aout 2026",
         f"({dc.dossier_du_mois(2026, 8, faux)})")
verifier("et malgré une faute de frappe",
         dc.dossier_du_mois(2026, 7, faux) == faux / "2026" / "juuillet 2026",
         f"({dc.dossier_du_mois(2026, 7, faux)})")
verifier("un dossier de factures n'est pas pris pour celui du mois",
         dc.dossier_du_mois(2026, 5, faux) is None,
         f"({dc.dossier_du_mois(2026, 5, faux)})")
verifier("rien n'est créé sans qu'on le demande",
         not (faux / "2026" / "mai 2026").exists())
cree = dc.dossier_du_mois(2026, 5, faux, creer=True)
verifier("le dossier manquant se crée au nom canonique",
         cree == faux / "2026" / "mai 2026" and cree.is_dir(), f"({cree})")
verifier("une année absente se crée aussi",
         dc.dossier_du_mois(2027, 3, faux, creer=True)
         == faux / "2027" / "mars 2027")

lu = dc.lire(transmis_aout)
verifier("un fichier transmis se lit", lu is not None)
verifier("son mois vient de ses dates, pas de son nom",
         lu and lu["mois"] == "2026-08", f"({lu and lu['mois']})")
verifier("le CA à déclarer est retrouvé par son en-tête",
         lu and abs(lu["totaux"]["ca_declare"] - 31 * 500.0) < 0.5,
         f"({lu and lu['totaux'].get('ca_declare')})")
verifier("les patients aussi",
         lu and abs(lu["totaux"]["patients"] - 31 * 5) < 0.5,
         f"({lu and lu['totaux'].get('patients')})")
verifier("la TVA déduite aussi", lu and "tva_deduite" in lu["totaux"])
verifier("un poste absent du fichier n'est pas inventé",
         lu and "recettes" not in lu["totaux"], f"({lu and sorted(lu['totaux'])})")
verifier("une facture n'est pas lue comme un mois",
         dc.lire(faux / "2026" / "factures mai 2026" / "steg.xlsx") is None)

verifier("l'inventaire ne retient que les feuilles mensuelles",
         len(dc.inventaire(faux)) == 2, f"({len(dc.inventaire(faux))})")
verifier("il balaie bien les trois classeurs",
         len(dc.fichiers(faux)) == 3, f"({len(dc.fichiers(faux))})")
verifier("un bordereau à nous se reconnaît",
         dc.est_notre_bordereau("bordereau-2026-08-aout.xlsx")
         and not dc.est_notre_bordereau("aout 2026.xlsx"))

# le bordereau rejoint le dossier du mois sans rien recouvrir
avant_octets = transmis_aout.read_bytes()
chemin_b, _ = bordereau.enregistrer(contenu, 2026, 8, racine_comptable=faux)
verifier("le bordereau se range dans le dossier du mois",
         chemin_b.parent == faux / "2026" / "aout 2026", f"({chemin_b.parent})")
verifier("sous un nom qui ne ressemble à aucun autre",
         chemin_b.name == "bordereau-2026-08-aout.xlsx", f"({chemin_b.name})")
verifier("le fichier du comptable est intact",
         transmis_aout.read_bytes() == avant_octets)
verifier("les deux cohabitent",
         len(list((faux / "2026" / "aout 2026").glob("*.xlsx"))) == 2)
verifier("mais le nôtre reste hors de la lecture des fichiers transmis",
         chemin_b not in dc.fichiers(faux),
         f"({[p.name for p in dc.fichiers(faux)]})")

print("\n11 bis. Le cache de ce qui a été transmis")
dc_racine = dc.RACINE
rapport = store.rafraichir_transmis(conn, racine=faux)
verifier("le premier passage lit tout",
         rapport["nouveaux"] == 2 and rapport["inchanges"] == 0, str(rapport))
verifier("les fichiers sans feuille mensuelle sont comptés à part",
         rapport["ignores"] >= 1, str(rapport))
par_mois = store.transmis_par_mois(conn)
verifier("un mois est indexé par son mois", "2026-08" in par_mois,
         f"({sorted(par_mois)})")
verifier("ses totaux sont relus depuis le cache",
         abs(par_mois["2026-08"]["totaux"]["ca_declare"] - 15500.0) < 0.5,
         f"({par_mois['2026-08']['totaux'].get('ca_declare')})")

rapport = store.rafraichir_transmis(conn, racine=faux)
verifier("le second passage ne rouvre rien",
         rapport["nouveaux"] == 0 and rapport["relus"] == 0
         and rapport["inchanges"] == rapport["vus"], str(rapport))
verifier("et le bordereau qu'on vient d'écrire n'est pas du lot",
         rapport["vus"] == 3, str(rapport))

fabriquer_transmis(transmis_aout, 2026, 8, 31, 600.0, 6)
rapport = store.rafraichir_transmis(conn, racine=faux)
verifier("un fichier modifié est relu", rapport["relus"] == 1, str(rapport))
verifier("et ses nouveaux totaux prennent la place",
         abs(store.transmis_par_mois(conn)["2026-08"]["totaux"]["ca_declare"]
             - 18600.0) < 0.5)

transmis_aout.unlink()
rapport = store.rafraichir_transmis(conn, racine=faux)
verifier("un fichier disparu quitte le cache", rapport["disparus"] == 1,
         str(rapport))
verifier("son mois n'est plus au tableau",
         "2026-08" not in store.transmis_par_mois(conn),
         f"({sorted(store.transmis_par_mois(conn))})")
verifier("le dossier réel du comptable n'a pas bougé de place",
         dc.RACINE == dc_racine, f"({dc.RACINE})")

print("\n11 ter. Les mois tracés sur la courbe d'évolution")
# Le classeur est pré-rempli à zéro jusqu'au 23/01/2027 : ces mois feraient
# plonger les courbes, et la chute se lirait comme un effondrement.
agregats_essai = [
    {"mois": "2026-06", "recettes": 30000.0},
    {"mois": "2026-07", "recettes": 35000.0},
    {"mois": "2026-08", "recettes": 31955.0},
    {"mois": "2026-09", "recettes": 0.0},
    {"mois": "2026-10", "recettes": 0.0},
    {"mois": "2026-11", "recettes": None},
]
traces = bordereau.mois_a_tracer(agregats_essai)
verifier("les mois sans recette sortent de la courbe",
         [r["mois"] for r in traces] == ["2026-06", "2026-07", "2026-08"],
         f"({[r['mois'] for r in traces]})")
verifier("un montant absent ne fait pas tomber la fonction",
         all(r["mois"] != "2026-11" for r in traces))
traces = bordereau.mois_a_tracer(agregats_essai, {"2026-10"})
verifier("un mois sans recette mais transmis reste tracé",
         "2026-10" in [r["mois"] for r in traces],
         f"({[r['mois'] for r in traces]})")
verifier("un mois vide et non transmis reste écarté",
         "2026-09" not in [r["mois"] for r in traces])
verifier("l'ordre des mois est conservé",
         [r["mois"] for r in traces] == sorted(r["mois"] for r in traces))

print("\n12. Le classeur du cabinet n'est jamais touché")
source = (PROJET / "compta" / "bordereau.py").read_text(encoding="utf-8")
verifier("le module ne connaît pas le chemin du classeur",
         "XLSX_PATH" not in source)
verifier("il n'ouvre aucun classeur existant",
         "load_workbook" not in source)
verifier("il ne fabrique que des classeurs neufs",
         source.count("Workbook()") == 1)

print(f"\n{'=' * 50}\n{ok} réussis, {ko} échecs\n{'=' * 50}")
sys.exit(1 if ko else 0)
