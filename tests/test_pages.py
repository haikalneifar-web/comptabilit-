"""Rendu des pages de l'app, sur une COPIE de la base (AppTest, sans navigateur)."""

import calendar
import datetime as dt
import shutil
import sys
import tempfile
from pathlib import Path

PROJET = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJET))

# le cabinet fictif de tests/atelier.py : ni le classeur réel, ni data/compta.db
import atelier  # noqa: E402

CLASSEUR, BASE = atelier.installer()

from compta import store, xlsx_io  # noqa: E402

tmp = Path(tempfile.mkdtemp()) / "test.db"
shutil.copy2(BASE, tmp)
store.DB_PATH = tmp

# une journée rectifiée, pour que les pages aient de la matière à afficher
conn = store.connect()
store.init_db(conn)
D = dt.date(2026, 7, 15)
store.ajouter_rectification(conn, D, "pay_n", store.lire_jour(conn, D).pay_n - 2,
                            "2 patients comptés deux fois (test)")
store.ajouter_rectification(conn, D, "divers", 120.0, "recette divers oubliée (test)")
conn.close()

from streamlit.testing.v1 import AppTest  # noqa: E402

ok, ko = 0, 0


def verifier(nom, condition, detail=""):
    global ok, ko
    if condition:
        ok += 1
        print(f"  OK   {nom}")
    else:
        ko += 1
        print(f"  ECHEC {nom} {detail}")


def table(at, *colonnes):
    """Le premier tableau qui porte ces colonnes — plus robuste qu'un index."""
    for d in at.dataframe:
        if all(c in d.value.columns for c in colonnes):
            return d.value
    return None


def metrique(at, libelle):
    """Une métrique retrouvée par son libellé, pas par sa position."""
    for m in at.metric:
        if m.label == libelle:
            return m
    return None


def nombre_de_tuile(texte):
    """Le nombre lu sur une tuile : « 1 724,43 TND » → 1724.43."""
    chiffres = "".join(c for c in str(texte) if c.isdigit() or c == ",")
    return float(chiffres.replace(",", ".")) if chiffres else 0.0


def lancer(page=None, etat=None):
    at = AppTest.from_file(str(PROJET / "streamlit_app.py"), default_timeout=30)
    for k, v in (etat or {}).items():
        at.session_state[k] = v
    at.run()
    if page:
        at.switch_page(page).run()
    return at


print("\n0. La page d'accueil est le mois comptable, ouvert sur le mois à saisir")
at = lancer()
verifier("aucune exception", not at.exception, str(at.exception))
verifier("le mois comptable est la page par défaut",
         any("Mois comptable" in h.value for h in at.header))
conn = store.connect()
attendu_mois = store.mois_a_saisir(conn)
dernier = store.dernier_jour_rempli(conn)
actifs_mois = conn.execute(
    "SELECT COUNT(*) n FROM jours WHERE substr(date,1,7)=? AND recettes>0",
    (f"{attendu_mois[0]:04d}-{attendu_mois[1]:02d}",),
).fetchone()["n"]
conn.close()

# Le mois proposé dépend de l'avancement de la saisie : on vérifie ses bornes,
# pas sa valeur, qui change à chaque mois repris.
aujourdhui = dt.date.today()
verifier("le mois proposé part de la dernière journée renseignée, sans dépasser "
         "le mois courant",
         (dernier.year, dernier.month) <= attendu_mois
         <= (aujourdhui.year, aujourdhui.month),
         f"({attendu_mois}, dernier={dernier})")
verifier("c'est bien lui qui est retenu à l'affichage",
         at.session_state["mois_travail"] == attendu_mois,
         f"({at.session_state['mois_travail']} au lieu de {attendu_mois})")
if actifs_mois:
    print(f"       (le mois proposé porte déjà {actifs_mois} journée(s) : "
          "pas d'avertissement attendu)")
else:
    verifier("le mois non saisi est signalé",
             any("n'est pas encore saisi" in w.value for w in at.warning))
verifier("toutes les journées du mois sont affichées",
         len(at.dataframe[0].value) == calendar.monthrange(*attendu_mois)[1],
         f"({len(at.dataframe[0].value)} pour {attendu_mois})")
verifier("la page d'accueil ne porte que la grille de saisie",
         len(at.dataframe) == 1, f"({len(at.dataframe)} tableaux)")

print("\n0 a. L'ordre des onglets")
# L'ordre de la barre du haut est un choix de l'utilisateur, pas un accident de
# rédaction. `st.navigation` ne l'expose pas à AppTest : on le relit à la source.
import re as _re_nav  # noqa: E402

_source_nav = (PROJET / "streamlit_app.py").read_text(encoding="utf-8")
# « Démarrage » n'est pas un onglet : c'est la page unique tant que le registre
# n'a pas été importé
onglets = [p for p in _re_nav.findall(r'st\.Page\("app_pages/(\w+)\.py"',
                                      _source_nav) if p != "demarrage"]
verifier("saisie, totaux, rectifications ouvrent la barre",
         onglets[:3] == ["journal", "totaux", "rectifications"],
         f"({onglets})")
verifier("les huit pages sont déclarées une fois chacune",
         len(onglets) == len(set(onglets)) == 8, f"({onglets})")

# Ce que portent les documents qui sortent de l'app se relit à la source : le
# sous-titre d'un export n'est pas observable depuis AppTest, qui n'expose pas
# le contenu d'un bouton de téléchargement. On ne regarde que les `sous_titre`,
# pas les légendes à l'écran, qui disent autre chose et n'accompagnent aucun
# document transmis.
import ast as _ast_exp  # noqa: E402


def _sous_titres_exports():
    trouves = []
    for chemin in sorted((PROJET / "app_pages").glob("*.py")):
        for noeud in _ast_exp.walk(_ast_exp.parse(chemin.read_text(encoding="utf-8"))):
            if not isinstance(noeud, _ast_exp.Call):
                continue
            for mot_cle in noeud.keywords:
                if mot_cle.arg != "sous_titre":
                    continue
                try:
                    trouves.append(_ast_exp.literal_eval(mot_cle.value))
                except (ValueError, SyntaxError):
                    # f-string : on ne garde que ses morceaux littéraux
                    trouves.append("".join(
                        part.value
                        for part in getattr(mot_cle.value, "values", [])
                        if isinstance(part, _ast_exp.Constant)))
    return trouves


exports = _sous_titres_exports()
verifier("chaque export porte un sous-titre", len(exports) >= 3, f"({exports})")
fautifs = [s for s in exports if "valeurs d'origine plus les" in s]
verifier("aucun ne porte « valeurs d'origine plus les rectifications motivées »",
         not fautifs, f"({fautifs})")

print("\n0 bis. Un mois au-delà de la dernière ligne du classeur")
at = lancer(etat={"mois_travail": (2027, 3)})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("mars 2027 s'affiche entièrement", len(at.dataframe[0].value) == 31,
         f"({len(at.dataframe[0].value)})")
verifier("les journées absentes du registre sont signalées",
         any("absente(s) du registre" in c.value for c in at.caption),
         str([c.value for c in at.caption]))

print("\n0 quater. Un outil d'export sous chaque tableau")
at = lancer(etat={"mois_travail": (2026, 7)})
boutons = at.get("download_button")
verifier("trois formats sous la grille de saisie",
         [b.label for b in boutons] == ["PDF", "Excel", "CSV"],
         f"({[b.label for b in boutons]})")
verifier("les boutons sont rattachés au mois affiché",
         all("2026_7" in b.key for b in boutons), f"({[b.key for b in boutons]})")
verifier("c'est bien l'export de la saisie",
         {b.key.split("_")[1] for b in boutons} == {"saisie"},
         f"({[b.key for b in boutons]})")

at = lancer("app_pages/totaux.py", {"mois_travail": (2026, 7)})
boutons = at.get("download_button")
verifier("trois formats sous le tableau des totaux",
         [b.label for b in boutons] == ["PDF", "Excel", "CSV"],
         f"({[b.label for b in boutons]})")
verifier("c'est bien l'export des colonnes calculées",
         {b.key.split("_")[1] for b in boutons} == {"calcul"},
         f"({[b.key for b in boutons]})")

at = lancer("app_pages/rectifications.py", {"mois_travail": (2026, 7)})
boutons = at.get("download_button")
verifier("deux tableaux exportables sur la page des rectifications",
         len(boutons) == 6, f"({len(boutons)})")
verifier("le corrigé et le journal sont proposés",
         {b.key.split("_")[1] for b in boutons} == {"corrige", "journal"},
         f"({[b.key for b in boutons]})")
verifier("chacun dans les trois formats",
         sorted(b.label for b in boutons) == sorted(["PDF", "Excel", "CSV"] * 2))
verifier("le registre réel ne s'exporte plus d'ici : il est sur Mois comptable",
         not any("reel" in b.key for b in boutons),
         f"({[b.key for b in boutons]})")

print("\n0 ter. Horizon des années et navigation de mois en mois")
at = lancer(etat={"mois_travail": (2026, 12)})
annees = [s for s in at.selectbox if s.label == "Année"][0].options
verifier("le choix des années remonte au début du registre", "2015" in annees)
verifier("il ne s'arrête pas à la dernière ligne du classeur",
         "2028" in annees and "2031" in annees, f"(jusqu'à {annees[-1]})")
verifier("aucune année manquante",
         len(annees) == int(annees[-1]) - int(annees[0]) + 1)

at.button(key="mois_suivant").click().run()
verifier("le mois suivant franchit l'année",
         at.session_state["mois_travail"] == (2027, 1),
         f"({at.session_state['mois_travail']})")
at.button(key="mois_suivant").click().run()
at.button(key="mois_precedent").click().run()
at.button(key="mois_precedent").click().run()
verifier("le mois précédent revient au même point",
         at.session_state["mois_travail"] == (2026, 12),
         f"({at.session_state['mois_travail']})")
verifier("aucune exception en naviguant", not at.exception, str(at.exception))

at.switch_page("app_pages/rectifications.py").run()
verifier("le mois suit d'une page à l'autre",
         at.session_state["mois_travail"] == (2026, 12),
         f"({at.session_state['mois_travail']})")
verifier("la page des rectifications affiche le même mois",
         [s for s in at.selectbox if s.label == "Mois"][0].value in (12, "Décembre"),
         f"({[s for s in at.selectbox if s.label == 'Mois'][0].value})")

print("\n1. Saisie du jour — journée ouverte")
at = lancer("app_pages/saisie.py", {"date_saisie": dt.date(2027, 6, 1)})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("mode de saisie proposé", len(at.segmented_control) >= 1)
verifier("bouton de clôture présent",
         any("Clôturer" in b.label for b in at.button))

print("\n2. Saisie — journée clôturée et rectifiée")
at = lancer("app_pages/saisie.py", {"date_saisie": D})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("bandeau de clôture", any("clôturée" in i.value for i in at.info))
verifier("bouton de rectification",
         any("Rectifier" in b.label for b in at.button))
textes = [m.value for m in at.markdown] + [c.value for c in at.caption]
verifier("les rectifications sont listées",
         any("comptés deux fois" in t for t in textes))

print("\n3. Saisie — la rectification est bien celle affichée")
at = lancer("app_pages/saisie.py", {"date_saisie": D})
conn = store.connect()
corrige = store.lire_jour_corrige(conn, D)
conn.close()
verifier("recettes du bandeau = registre corrigé",
         at.metric[0].value == f"{corrige.recettes:,.0f}".replace(",", " ") + " TND",
         f"({at.metric[0].value})")

print("\n4. Mois comptable — la grille de saisie, et rien d'autre")
at = lancer("app_pages/journal.py", {"mois_travail": (2026, 7)})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("mention des journées rectifiées",
         any("rectifiée" in c.value for c in at.caption))
verifier("aucune modification en attente au chargement",
         any("Aucune." in c.value for c in at.caption))

verifier("un seul tableau : la grille de saisie", len(at.dataframe) == 1,
         f"({len(at.dataframe)})")
saisie_df = at.dataframe[0].value
verifier("la grille porte les colonnes de saisie",
         "pay_n" in saisie_df.columns and "divers" in saisie_df.columns)
verifier("aucune colonne calculée ne s'y est glissée",
         not ({"Recettes", "CA déclaré", "Caisse", "Dû CNAM"}
              & set(saisie_df.columns)), f"({list(saisie_df.columns)})")
verifier("la page renvoie vers les totaux",
         any("formules en tirent" in p.label for p in at.get("page_link")),
         f"({[p.label for p in at.get('page_link')]})")

# --- mise en page ------------------------------------------------------------
# Les en-têtes étaient rognés — « Rembo », « Catarac », « Laser d » — parce que
# douze colonnes se partageaient la largeur automatiquement. Les largeurs sont
# désormais posées en pixels, et ce test vérifie qu'elles tiennent dans la page.
# Ajouter une colonne sans rééquilibrer le fera échouer.
import json as _json_mep  # noqa: E402

LARGEUR_PAGE = 1324               # fenêtre de l'app, marges déduites


def largeurs(tableau):
    reglages = _json_mep.loads(tableau.proto.columns or "{}")
    return {cle: r.get("width") for cle, r in reglages.items()
            if isinstance(r, dict) and r.get("width")}


def tient_dans_la_page(at_page, nom, marge=1.0):
    for rang_tab, tableau in enumerate(at_page.dataframe):
        posees = largeurs(tableau)
        affichees = set(tableau.value.columns) | {"Date"}
        manquantes = affichees - set(posees)
        verifier(f"{nom} : chaque colonne affichée a une largeur posée",
                 not manquantes, f"(sans largeur : {sorted(manquantes)})")
        demande = sum(posees[c] for c in affichees if c in posees)
        verifier(f"{nom} : ses colonnes tiennent dans la largeur de page",
                 demande <= LARGEUR_PAGE * marge,
                 f"({demande} px demandés pour "
                 f"{LARGEUR_PAGE * marge:.0f} disponibles)")


def parcours(bloc, volet=None):
    """L'arbre de rendu à plat : (type, libellé, volet qui le contient)."""
    trouves = []
    for enfant in getattr(bloc, "children", {}).values():
        genre = getattr(enfant, "type", enfant.__class__.__name__)
        try:
            etiquette = enfant.label
        except Exception:                      # noqa: BLE001 — sans libellé
            etiquette = None
        trouves.append((genre, etiquette, volet))
        trouves += parcours(enfant, etiquette if genre == "status" else volet)
    return trouves


bandeau = [m.label for m in at.metric]
verifier("un bandeau de six totaux, sous le tableau de saisie",
         bandeau == ["Patients", "Chirurgie", "Recettes", "CA déclaré",
                     "Caisse", "Dû CNAM"], f"({bandeau})")

# L'ordre d'affichage se lit dans l'arbre de rendu, seul endroit où il existe :
# `at.dataframe` et `at.metric` rendent les éléments à plat, sans dire lequel
# vient avant. C'est la demande — les tuiles SOUS le tableau — qu'on vérifie ici.
arbre = parcours(at.main)
position = lambda test: next((i for i, n in enumerate(arbre) if test(n)), None)  # noqa: E731
i_grille = position(lambda n: n[0] == "dataframe")
i_tuiles = position(lambda n: n[0] == "metric")
verifier("les tuiles de totaux viennent sous le tableau",
         i_grille is not None and i_tuiles is not None and i_grille < i_tuiles,
         f"(grille en {i_grille}, tuiles en {i_tuiles})")

tient_dans_la_page(at, "grille de saisie")
etiquettes = [
    r.get("label") for r
    in _json_mep.loads(at.dataframe[0].proto.columns or "{}").values()
    if isinstance(r, dict) and r.get("label")
]
verifier("les libellés de la grille tiennent dans leur colonne",
         all(len(str(e)) <= 13 for e in etiquettes), str(etiquettes))

at.segmented_control(key="registre_journal").select("Réel").run()
verifier("bascule vers le réel sans exception", not at.exception, str(at.exception))
verifier("valeurs d'origine annoncées",
         any("avant rectification" in c.value for c in at.caption))
verifier("registre réel non éditable",
         any("ne s'éditent pas" in i.value for i in at.info))

print("\n4 a. Mois comptable — le tarif se lit au-dessus de la grille")
# La règle est de régler le tarif AVANT de saisir, une journée enregistrée
# figeant le sien. L'écran de saisie ne le montrait nulle part. On pose un tarif
# APCI que rien d'autre ne porte, pour le reconnaître à coup sûr.
conn = store.connect()
tarifs_avant = store.tarifs_courants(conn)
store.enregistrer_tarifs(conn, tarifs_avant | {"t_apci": 99.0})
a_saisir = store.mois_a_saisir(conn)
debut_a_saisir = f"{a_saisir[0]:04d}-{a_saisir[1]:02d}"
apci_classeur = sorted({r["t_apci"] for r in conn.execute(
    "SELECT t_apci FROM jours WHERE substr(date,1,7)=? AND cloture=0 AND "
    "(pay_n+remb_n+apci_n+tp_n+star_n+cat_cnam_n+cat_pay_n)=0",
    (debut_a_saisir,))})
conn.close()


def legende_tarifs(at_page):
    return next((c.value for c in at_page.caption
                 if "Tarifs de ce mois" in c.value), None)


def a_plat(bloc):
    """L'arbre de rendu à plat, chaque nœud tel quel, dans l'ordre d'affichage."""
    noeuds = []
    for enfant in getattr(bloc, "children", {}).values():
        noeuds.append(enfant)
        noeuds += a_plat(enfant)
    return noeuds


at = lancer("app_pages/journal.py", {"mois_travail": a_saisir})
verifier("aucune exception", not at.exception, str(at.exception))
legende = legende_tarifs(at)
verifier("le mois à saisir annonce ses tarifs", legende is not None)
verifier("au tarif en vigueur, pas à celui du classeur",
         legende is not None and "APCI" in legende and "**99**" in legende,
         f"({legende})")
if apci_classeur:
    verifier("le prix encore inscrit au classeur est montré barré",
             f"~~{'/'.join(f'{v:g}' for v in apci_classeur)}~~" in (legende or ""),
             f"({legende}, classeur {apci_classeur})")
verifier("et un lien mène aux tarifs, pour les régler avant de saisir",
         any(p.label == "Régler les tarifs" for p in at.get("page_link")))
noeuds = a_plat(at.main)
i_legende = next((i for i, n in enumerate(noeuds)
                  if "Tarifs de ce mois" in str(getattr(n, "value", ""))), None)
i_grille = next((i for i, n in enumerate(noeuds)
                 if getattr(n, "type", "") in ("dataframe", "arrow_data_frame")), None)
verifier("la légende vient avant la grille",
         i_legende is not None and i_grille is not None and i_legende < i_grille,
         f"(légende en {i_legende}, grille en {i_grille})")

at = lancer("app_pages/journal.py", {"mois_travail": (2026, 7)})
legende = legende_tarifs(at)
verifier("un mois déjà travaillé garde ses tarifs, pas celui du jour",
         legende is not None and "99" not in legende and "déjà travaillé" in legende,
         f"({legende})")
verifier("et n'offre pas de les régler",
         not any(p.label == "Régler les tarifs" for p in at.get("page_link")))

conn = store.connect()
store.enregistrer_tarifs(conn, tarifs_avant)
conn.close()

print("\n4 bis. Totaux du mois — les colonnes calculées, sur leur page")
at = lancer("app_pages/totaux.py", {"mois_travail": (2026, 7)})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("la page porte son titre",
         any("Totaux du mois" in h.value for h in at.header))
# Un seul tableau : ce que le classeur tire du mois. Les montants par poste ne
# sont qu'en tuiles — leur détail journée par journée reprendrait la grille de
# saisie de la page précédente, au tarif près.
verifier("un seul tableau : ce que le classeur en tire",
         len(at.dataframe) == 1, f"({len(at.dataframe)})")
calcul = at.dataframe[0].value

# Le garde-fou : la page doit montrer TOUTES les colonnes que le classeur
# calcule. Onze manquaient — dont « à verser à la banque » (AA), la TVA déduite
# (AD) et les neuf montants par poste — et rien ne le signalait.
from compta.config import CALC_COLS  # noqa: E402
from compta.model import CALCULES  # noqa: E402

verifier("la liste des colonnes calculées couvre celles du classeur",
         {(attribut, lettre) for lettre, _, attribut, _ in CALCULES}
         == set(CALC_COLS.items()),
         f"(oubliées : {sorted(set(CALC_COLS.values()) - {l for l, _, _, _ in CALCULES})})")
a_l_ecran = set(calcul.columns) | {m.label for m in at.metric}
manquantes = [f"{lettre} ({libelle})" for lettre, libelle, _, _ in CALCULES
              if libelle not in a_l_ecran]
verifier("et les seize sont à l'écran, en tuile ou en colonne",
         not manquantes, f"(absentes de l'écran : {manquantes})")
verifier("« à verser à la banque » en fait partie",
         "À verser banque" in a_l_ecran, f"({sorted(a_l_ecran)})")
verifier("les montants par poste ne sont qu'en tuiles",
         all(libelle not in calcul.columns
             for _, libelle, _, groupe in CALCULES if groupe == "poste"),
         f"({[l for _, l, _, g in CALCULES if g == 'poste' and l in calcul.columns]})")
verifier("le tableau porte ce que le classeur en tire",
         all(libelle in calcul.columns
             for _, libelle, _, groupe in CALCULES if groupe == "total"),
         f"({list(calcul.columns)})")

# L'export, lui, garde les seize colonnes : c'est le document qu'on sort.
# AppTest n'expose pas le contenu d'un bouton de téléchargement — on vérifie
# donc la source, le calcul dont l'export est tiré.
verifier("les trois formats sont proposés",
         sorted(b.label for b in at.get("download_button"))
         == ["CSV", "Excel", "PDF"])
styles_totaux = "\n".join(m.value for m in at.markdown)
verifier("les deux tuiles colorées le sont aussi sur cette page",
         ".st-key-tuile_banque" in styles_totaux
         and ".st-key-tuile_ca_declare" in styles_totaux,
         "(pas de teinte sur la page des totaux)")
from app_utils import calculs_du_mois as _calculs  # noqa: E402
from app_utils import mois_affiche as _mois  # noqa: E402

conn = store.connect()
_tarifs = store.tarifs_courants(conn)
_, lignes_calculees = _calculs(_mois(conn, 2026, 7, True, _tarifs), _tarifs,
                               store.frontiere_tarifs(conn))
conn.close()
oubliees = [libelle for _, libelle, _, _ in CALCULES
            if libelle not in lignes_calculees[0]]
verifier("le calcul d'où sort l'export porte les seize colonnes",
         not oubliees, f"(absentes : {oubliees})")
verifier("elle dit que rien ne s'y saisit",
         any("Rien ne se saisit" in c.value for c in at.caption))
verifier("elle renvoie vers la saisie",
         any("Retour à la saisie" in p.label for p in at.get("page_link")),
         f"({[p.label for p in at.get('page_link')]})")
tient_dans_la_page(at, "tableau des totaux")

conn = store.connect()
attendu = {d: store.lire_jour_corrige(conn, d)
           for d in store.jours_ouverts(conn, dt.date(2026, 7, 1), dt.date(2026, 7, 31))}
jour_corrige = store.lire_jour_corrige(conn, D)
ouvertes = list(attendu)
conn.close()

verifier("une ligne par journée du mois", len(calcul) == 31, f"({len(calcul)})")
verifier("recettes calculées = registre corrigé",
         abs(float(calcul.loc[D, "Recettes"]) - jour_corrige.recettes) < 0.01,
         f"({calcul.loc[D, 'Recettes']} vs {jour_corrige.recettes})")
verifier("patients calculés = registre corrigé",
         int(calcul.loc[D, "Patients"]) == jour_corrige.patients)
verifier("rien d'autre que du calculé n'y figure",
         set(calcul.columns)
         <= {libelle for _, libelle, _, _ in CALCULES} | {"Δ en attente"},
         f"(en trop : {sorted(set(calcul.columns) - {l for _, l, _, _ in CALCULES})})")
verifier("l'écart en attente reste caché sans saisie en attente",
         "Δ en attente" not in calcul.columns)
verifier("la journée rectifiée n'est pas un emplacement ouvert", D not in ouvertes)

# --- les tuiles portent les totaux du mois des colonnes calculées ------------
tuiles = {m.label: m for m in at.metric}
verifier("une tuile par colonne calculée, les seize",
         list(tuiles) == [libelle for _, libelle, _, _ in CALCULES],
         f"({list(tuiles)})")
conn = store.connect()
journees_juillet = [store.lire_jour_corrige(conn, dt.date(2026, 7, j),
                                            store.tarifs_courants(conn))
                    for j in range(1, 32)]
conn.close()
for lettre, libelle, attribut, _ in CALCULES:
    attendu_total = sum(getattr(j, attribut) for j in journees_juillet)
    verifier(f"tuile {lettre} ({libelle}) = total du mois",
             abs(nombre_de_tuile(tuiles[libelle].value) - abs(attendu_total)) <= 1.0,
             f"(tuile {tuiles[libelle].value!r}, registre {attendu_total:.2f})")

at.segmented_control(key="registre_totaux").select("Réel").run()
verifier("bascule vers le réel sans exception", not at.exception, str(at.exception))
calcul_reel = at.dataframe[0].value
verifier("le réel diffère du corrigé sur la journée rectifiée",
         abs(float(calcul_reel.loc[D, "Recettes"])
             - float(calcul.loc[D, "Recettes"])) > 0.01)
verifier("les autres journées sont identiques dans les deux registres",
         all(abs(float(calcul_reel.loc[d, "Recettes"])
                 - float(calcul.loc[d, "Recettes"])) < 0.01
             for d in calcul.index if d != D))

print("\n5. Mois comptable — un mois hors du classeur, entièrement ouvert")
conn = store.connect()
from compta.model import Journee  # noqa: E402
j = Journee(date=dt.date(2027, 6, 3), pay_n=5, source="app")
j.tarifs = store.tarifs_courants(conn)
store.ecrire_jour(conn, j)
attendu_juin = store.lire_jour(conn, dt.date(2027, 6, 3))
conn.close()
at = lancer("app_pages/journal.py", {"mois_travail": (2027, 6)})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("la grille porte les 30 journées de juin", len(at.dataframe[0].value) == 30,
         f"({len(at.dataframe[0].value)})")
verifier("aucun motif demandé sans modification",
         not any("Motif des rectifications" in t.label for t in at.text_area))

at = lancer("app_pages/totaux.py", {"mois_travail": (2027, 6)})
verifier("les totaux de ce mois s'affichent aussi", not at.exception,
         str(at.exception))
calcul = at.dataframe[0].value
verifier("ses recettes sont calculées",
         abs(float(calcul.loc[dt.date(2027, 6, 3), "Recettes"])
             - attendu_juin.recettes) < 0.01)

print("\n5 bis. Mois comptable — modifier une cellule, puis enregistrer")
# AppTest ne clique pas dans une grille : on injecte l'état du widget, ce que
# `st.data_editor` relit comme des modifications en attente.
GRILLE = "grille_2026_7"
POS_15 = 14  # 15 juillet, 15e ligne du mois
conn = store.connect()
avant = store.lire_jour_corrige(conn, D)
reel_avant = store.lire_jour(conn, D)
nb_rectifs_avant = store.nb_rectifications(conn)
conn.close()

at = lancer("app_pages/journal.py", {
    "mois_travail": (2026, 7),
    GRILLE: {"edited_rows": {POS_15: {"pay_n": avant.pay_n + 5}},
             "added_rows": [], "deleted_rows": []},
})
verifier("aucune exception", not at.exception, str(at.exception))
attendu_ecart = 5 * avant.tarifs["t_pay"]
tuile = metrique(at, "Recettes")
verifier("la tuile des recettes chiffre la modification en attente",
         tuile is not None and tuile.delta
         and tuile.delta.replace(" ", "").startswith(f"+{attendu_ecart:.0f}"),
         f"({tuile.delta if tuile else None})")

# la saisie en attente vit dans l'état de session : elle suit sur l'autre page
at_totaux = at.switch_page("app_pages/totaux.py").run()
verifier("la page des totaux reprend la saisie en attente",
         not at_totaux.exception, str(at_totaux.exception))
calcul = at_totaux.dataframe[0].value
verifier("la colonne calculée suit la modification",
         abs(float(calcul.loc[D, "Δ en attente"]) - attendu_ecart) < 0.01,
         f"({calcul.loc[D, 'Δ en attente']})")
verifier("les recettes affichées intègrent la modification",
         abs(float(calcul.loc[D, "Recettes"])
             - (avant.recettes + attendu_ecart)) < 0.01)
verifier("les autres journées ne bougent pas",
         all(abs(float(calcul.loc[d, "Δ en attente"])) < 0.01
             for d in calcul.index if d != D))
verifier("la page le dit en toutes lettres",
         any("modification(s) en attente" in i.value for i in at_totaux.info),
         str([i.value for i in at_totaux.info]))

at = at.switch_page("app_pages/journal.py").run()
verifier("retour à la saisie sans exception", not at.exception, str(at.exception))
verifier("la modification est listée",
         any("Modifications en attente" in s.value for s in at.subheader))
verifier("un motif est exigé",
         any("Motif des rectifications" in t.label for t in at.text_area))
bouton = [b for b in at.button if "Enregistrer" in b.label][0]
verifier("enregistrement bloqué sans motif", bouton.disabled)

at.text_area(key="motif_2026_7").set_value("5 patients oubliés au report (test)").run()
bouton = [b for b in at.button if "Enregistrer" in b.label][0]
verifier("enregistrement débloqué avec motif", not bouton.disabled)
bouton.click().run()
verifier("aucune exception à l'enregistrement", not at.exception, str(at.exception))

conn = store.connect()
verifier("une rectification de plus au journal",
         store.nb_rectifications(conn) == nb_rectifs_avant + 1,
         f"({store.nb_rectifications(conn)})")
apres = store.lire_jour_corrige(conn, D)
verifier("registre corrigé à jour", apres.pay_n == avant.pay_n + 5)
verifier("registre réel intact", store.lire_jour(conn, D).pay_n == reel_avant.pay_n)
derniere_rectif = store.rectifications_du_jour(conn, D)[0]
verifier("le motif saisi est bien attaché",
         derniere_rectif["motif"] == "5 patients oubliés au report (test)")
verifier("l'ancienne valeur est celle d'avant",
         derniere_rectif["ancienne"] == avant.pay_n)
conn.close()

print("\n5 ter. Mois comptable — une journée ouverte s'enregistre sans motif")
conn = store.connect()
nb_rectifs_avant = store.nb_rectifications(conn)
ouverte = store.jours_ouverts(conn, dt.date(2026, 7, 1), dt.date(2026, 7, 31))[0]
conn.close()
at = lancer("app_pages/journal.py", {
    "mois_travail": (2026, 7),
    GRILLE: {"edited_rows": {ouverte.day - 1: {"pay_n": 3}},
             "added_rows": [], "deleted_rows": []},
})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("aucun motif demandé pour une journée ouverte",
         not any("Motif des rectifications" in t.label for t in at.text_area))
bouton = [b for b in at.button if "Enregistrer" in b.label][0]
verifier("enregistrement possible d'emblée", not bouton.disabled)
bouton.click().run()
verifier("aucune exception à l'enregistrement", not at.exception, str(at.exception))
conn = store.connect()
verifier("écriture directe, sans rectification",
         store.nb_rectifications(conn) == nb_rectifs_avant)
verifier("la journée ouverte porte la nouvelle valeur",
         store.lire_jour(conn, ouverte).pay_n == 3)
conn.close()

print("\n5 ter bis. Mois comptable — « Enregistrer dans le classeur » est à portée")
# l'écriture se fait sur une COPIE du classeur : l'original est partagé par tous les tests
copie_classeur = Path(tempfile.mkdtemp()) / "classeur.xlsx"
shutil.copy2(CLASSEUR, copie_classeur)
conn = store.connect()
xlsx_io.lier_classeur(conn, copie_classeur)
conn.close()
at = lancer("app_pages/journal.py", {"mois_travail": (2026, 7)})
verifier("aucune exception", not at.exception, str(at.exception))
bouton_classeur = at.button(key="classeur_mois")
verifier("le bouton classeur annonce les journées en attente",
         "(" in bouton_classeur.label, bouton_classeur.label)
verifier("il est actif quand rien n'est en attente de saisie",
         not bouton_classeur.disabled)
bouton_classeur.click().run()
verifier("aucune exception à l'écriture", not at.exception, str(at.exception))
conn = store.connect()
verifier("la journée est écrite et clôturée",
         ouverte not in store.jours_a_exporter(conn)
         and store.lire_jour(conn, ouverte).cloture)
conn.close()
verifier("le bouton est grisé une fois le classeur à jour",
         at.button(key="classeur_mois").disabled)

print("\n5 quater. Une saisie non enregistrée survit au changement d'onglet")
conn = store.connect()
avant_q = store.lire_jour_corrige(conn, D)
nb_rectifs_avant = store.nb_rectifications(conn)
conn.close()

at = lancer("app_pages/journal.py", {
    "mois_travail": (2026, 7),
    GRILLE: {"edited_rows": {POS_15: {"pay_n": avant_q.pay_n + 7}},
             "added_rows": [], "deleted_rows": []},
})
verifier("modification visible avant de changer d'onglet",
         any("Modifications en attente" in s.value for s in at.subheader)
         and (t := metrique(at, "Recettes")) is not None and t.delta)
verifier("le brouillon est mémorisé hors de l'état des widgets",
         at.session_state["_brouillons"][GRILLE].get((POS_15, "pay_n"))
         == avant_q.pay_n + 7,
         f"({at.session_state['_brouillons'].get(GRILLE)})")

# Aller-retour, en reproduisant ce que fait Streamlit dans le navigateur :
# l'état du widget est purgé quand la page n'est plus affichée. Seul le
# brouillon, clé ordinaire, doit permettre de retrouver la saisie.
at.switch_page("app_pages/tableau_bord.py").run()
verifier("passage sur un autre onglet sans exception", not at.exception,
         str(at.exception))
del at.session_state[GRILLE]
verifier("état du widget bien purgé", GRILLE not in at.session_state)
verifier("le brouillon a survécu à la purge",
         GRILLE in at.session_state["_brouillons"])

at.switch_page("app_pages/journal.py").run()
verifier("retour sans exception", not at.exception, str(at.exception))
attendu_q = 7 * avant_q.tarifs["t_pay"]
tuile_q = metrique(at, "Recettes")
verifier("la modification est restaurée depuis le brouillon",
         tuile_q is not None and tuile_q.delta
         and tuile_q.delta.replace(" ", "").startswith(f"+{attendu_q:.0f}"),
         f"({tuile_q.delta if tuile_q else None})")
verifier("et la grille la porte de nouveau",
         abs(float(at.dataframe[0].value.iloc[POS_15]["pay_n"])
             - (avant_q.pay_n + 7)) < 0.01,
         f"({at.dataframe[0].value.iloc[POS_15]['pay_n']})")
verifier("elle est signalée comme non enregistrée",
         any("Rien n'est encore enregistré" in c.value for c in at.caption))
verifier("rien n'a été enregistré entre-temps",
         store.lire_jour_corrige(store.connect(), D).pay_n == avant_q.pay_n)

# le motif aussi
at.text_area(key="motif_2026_7").set_value("motif en cours de frappe (test)").run()
at.switch_page("app_pages/rectifications.py").run()
at.switch_page("app_pages/journal.py").run()
verifier("le motif a survécu au détour",
         at.text_area(key="motif_2026_7").value == "motif en cours de frappe (test)",
         f"({at.text_area(key='motif_2026_7').value!r})")

# abandon explicite
at.button(key="abandon_mois").click().run()
verifier("après abandon, plus aucune modification",
         any("Aucune." in c.value for c in at.caption))
verifier("les tuiles sont revenues au registre",
         (t := metrique(at, "Recettes")) is not None and not t.delta,
         f"({t.delta if t else None})")
at.switch_page("app_pages/totaux.py").run()
verifier("et la page des totaux aussi : plus d'écart à montrer",
         "Δ en attente" not in at.dataframe[0].value.columns,
         f"({list(at.dataframe[0].value.columns)})")
verifier("ses tuiles sont revenues au registre",
         all(not m.delta for m in at.metric),
         f"({[(m.label, m.delta) for m in at.metric if m.delta]})")
at.switch_page("app_pages/journal.py").run()
verifier("l'abandon tient après un changement d'onglet",
         (t := metrique(at, "Recettes")) is not None and not t.delta,
         f"({t.delta if t else None})")
verifier("toujours rien en base",
         store.nb_rectifications(store.connect()) == nb_rectifs_avant)

print("\n6. Rectifications — le mois de juillet 2026")
conn = store.connect()
du_mois = store.liste_rectifications(conn, dt.date(2026, 7, 1), dt.date(2026, 7, 31))
attendu_total = len(du_mois)
attendu_vigueur = sum(1 for r in du_mois if r["en_vigueur"])
conn.close()

at = lancer("app_pages/rectifications.py", {"mois_travail": (2026, 7)})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("compteur de rectifications du mois",
         metrique(at, "Rectifications du mois").value == str(attendu_total),
         f"({metrique(at, 'Rectifications du mois')} vs {attendu_total})")
verifier("compteur des rectifications en vigueur",
         metrique(at, "En vigueur").value == str(attendu_vigueur),
         f"({metrique(at, 'En vigueur')} vs {attendu_vigueur})")
journal_df = table(at, "Motif", "En vigueur")
verifier("journal affiché", journal_df is not None)
verifier("les corrections du mois sont au journal",
         len(journal_df) == attendu_total, f"({len(journal_df)})")
verifier("toutes portent sur juillet 2026",
         all(d.month == 7 and d.year == 2026 for d in journal_df["Journée"]))
verifier("chaque motif de test est présent",
         all("test" in str(m) for m in journal_df["Motif"]))
verifier("une rectification remplacée n'est plus en vigueur",
         int(journal_df["En vigueur"].sum()) == attendu_vigueur)
comparaison = table(at, "Recettes réelles", "Recettes corrigées")
verifier("comparaison jour par jour", comparaison is not None)
verifier("la comparaison ne porte que sur les journées rectifiées",
         len(comparaison) == 1 and comparaison["Journée"].iloc[0] == D,
         f"({len(comparaison)} lignes)")
verifier("réel et corrigé diffèrent sur cette journée",
         abs(float(comparaison["Recettes réelles"].iloc[0])
             - float(comparaison["Recettes corrigées"].iloc[0])) > 0.01)
verifier("les mois porteurs de corrections sont listés",
         any("Mois avec rectifications" in c.value for c in at.caption))

print("\n6 bis. Rectifications — un mois sans aucune correction")
at = lancer("app_pages/rectifications.py", {"mois_travail": (2026, 6)})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("le mois est annoncé identique au réel",
         any("identique au registre réel" in i.value for i in at.info))
verifier("le cumul général reste visible",
         any("Sur l'ensemble du registre" in c.value for c in at.caption))

print("\n7. Rectifications — le tableau réel / corrigé")
conn = store.connect()
cloturees = [dt.date.fromisoformat(r["date"])
             for r in store.journal(conn, dt.date(2026, 7, 1), dt.date(2026, 7, 31),
                                    corrige=True) if r["cloture"]]
position = cloturees.index(D)
reel_D = store.lire_jour(conn, D)
corrige_D = store.lire_jour_corrige(conn, D)
nb_avant = store.nb_rectifications(conn)
conn.close()

at = lancer("app_pages/rectifications.py", {"mois_travail": (2026, 7)})
# Une seule grille désormais : le registre réel se lit sur « Mois comptable »
# en vue « Réel », il faisait double emploi ici.
corrige_df = at.dataframe[0].value
verifier("la grille montre le registre corrigé",
         int(corrige_df["pay_n"].iloc[position]) == corrige_D.pay_n,
         f"({corrige_df['pay_n'].iloc[position]} vs {corrige_D.pay_n})")
verifier("elle diffère du réel sur la journée déjà rectifiée",
         reel_D.pay_n != corrige_D.pay_n)
verifier("le registre réel n'a plus sa grille ici",
         not any(int(d.value["pay_n"].iloc[position]) == reel_D.pay_n
                 for d in at.dataframe if "pay_n" in d.value.columns),
         "(une grille porte encore les valeurs d'origine)")
verifier("seules les journées clôturées y figurent",
         len(corrige_df) == len(cloturees), f"({len(corrige_df)})")
verifier("aucune correction en attente au chargement",
         any("Aucune." in c.value for c in at.caption))
# la grille partage la configuration de colonnes de l'écran de saisie
grille_rectif = next((d for d in at.dataframe if "pay_n" in d.value.columns), None)
verifier("ses colonnes ont les largeurs posées de la grille de saisie",
         grille_rectif is not None
         and set(largeurs(grille_rectif)) >= set(corrige_df.columns) | {"Date"},
         f"({sorted(largeurs(grille_rectif)) if grille_rectif else None})")
verifier("et elles tiennent dans la page",
         grille_rectif is not None
         and sum(largeurs(grille_rectif).values()) <= LARGEUR_PAGE,
         f"({sum(largeurs(grille_rectif).values()) if grille_rectif else 0} px)")

conn = store.connect()
mois_reel = [store.vers_journee(r) for r in store.journal(
    conn, dt.date(2026, 7, 1), dt.date(2026, 7, 31)) if r["cloture"]]
conn.close()
recettes_reelles = sum(j.recettes for j in mois_reel)
verifier("un seul bandeau de totaux, celui du corrigé",
         sum(1 for m in at.metric if m.label == "Recettes") == 1,
         f"({sum(1 for m in at.metric if m.label == 'Recettes')})")
conn = store.connect()
mois_corrige = [store.vers_journee(r) for r in store.journal(
    conn, dt.date(2026, 7, 1), dt.date(2026, 7, 31), corrige=True) if r["cloture"]]
conn.close()
recettes_corrigees = sum(j.recettes for j in mois_corrige)
ecart_deja_la = recettes_corrigees - recettes_reelles
recettes = [m for m in at.metric if m.label == "Recettes"]
verifier("il porte les totaux du registre corrigé",
         recettes[0].value
         == f"{recettes_corrigees:,.0f}".replace(",", " ") + " TND",
         f"({recettes[0].value})")
verifier("et chiffre les rectifications déjà enregistrées, face au réel",
         recettes[0].delta.replace(" ", "").replace("TND", "")
         == f"{ecart_deja_la:+.0f}",
         f"({recettes[0].delta} vs {ecart_deja_la:+.0f})")

# --- les tuiles font le calcul sur les valeurs rectifiées -------------------
# Les mêmes seize colonnes que sur « Totaux du mois », mais calculées sur le
# registre corrigé, avec l'écart au réel sous chaque tuile.
tuiles_rectif = {m.label: m for m in at.metric}
# les cinq métriques du journal des corrections suivent les seize tuiles
verifier("une tuile par colonne calculée, sur les valeurs rectifiées",
         list(tuiles_rectif)[:16] == [libelle for _, libelle, _, _ in CALCULES],
         f"({list(tuiles_rectif)[:16]})")

# Le bandeau est au-dessus de la grille et y reste collé : on corrige une ligne
# et on voit bouger les totaux sans quitter la grille des yeux.
arbre_rectif = parcours(at.main)
place = lambda test: next((i for i, n in enumerate(arbre_rectif) if test(n)), None)  # noqa: E731
verifier("les tuiles passent au-dessus de la grille",
         (i_t := place(lambda n: n[0] == "metric")) is not None
         and (i_g := place(lambda n: n[0] == "dataframe")) is not None
         and i_t < i_g, f"(tuiles en {i_t}, grille en {i_g})")
styles = "\n".join(m.value for m in at.markdown)
verifier("et le bandeau est collé en haut", "position: sticky" in styles)
verifier("ses tuiles sont resserrées",
         "stMetricValue" in styles and "st-key-tuiles_rectif" in styles)

# deux colonnes se cherchent du regard : elles portent une couleur
from app_utils import COULEURS_TUILES  # noqa: E402

verifier("« à verser à la banque » est en jaune",
         f".st-key-tuile_banque" in styles
         and COULEURS_TUILES["banque"][0] in styles,
         f"({COULEURS_TUILES['banque']})")
verifier("« CA déclaré » est en vert clair",
         ".st-key-tuile_ca_declare" in styles
         and COULEURS_TUILES["ca_declare"][0] in styles,
         f"({COULEURS_TUILES['ca_declare']})")
verifier("seules ces deux-là sont teintées",
         set(COULEURS_TUILES) == {"banque", "ca_declare"},
         f"({sorted(COULEURS_TUILES)})")
verifier("et ce sont bien deux colonnes calculées du classeur",
         set(COULEURS_TUILES) <= {attribut for _, _, attribut, _ in CALCULES})
# seize tuiles sur deux rangées de huit — on ne compte que les colonnes situées
# avant la grille : celles du journal des corrections viennent bien après
colonnes_bandeau = sum(1 for n in arbre_rectif[:i_g] if n[0] == "column")
verifier("elles tiennent sur deux rangées de huit",
         colonnes_bandeau == 16, f"({colonnes_bandeau} colonnes avant la grille)")
for lettre, libelle, attribut, _ in CALCULES:
    attendu_corrige = sum(getattr(j, attribut) for j in mois_corrige)
    verifier(f"tuile {lettre} ({libelle}) = total rectifié du mois",
             abs(nombre_de_tuile(tuiles_rectif[libelle].value)
                 - abs(attendu_corrige)) <= 1.0,
             f"(tuile {tuiles_rectif[libelle].value!r}, "
             f"corrigé {attendu_corrige:.2f})")

ecarts_attendus = {
    libelle: sum(getattr(j, attribut) for j in mois_corrige)
             - sum(getattr(j, attribut) for j in mois_reel)
    for _, libelle, attribut, _ in CALCULES
}
touchees = [libelle for libelle, e in ecarts_attendus.items() if abs(e) > 0.005]
verifier("les rectifications déjà enregistrées touchent au moins une colonne",
         touchees, "(aucun écart réel → corrigé sur le mois)")
for libelle in touchees:
    verifier(f"la tuile {libelle} chiffre son écart au réel",
             tuiles_rectif[libelle].delta
             and abs(nombre_de_tuile(tuiles_rectif[libelle].delta)
                     - abs(ecarts_attendus[libelle])) <= 1.0,
             f"({tuiles_rectif[libelle].delta} vs "
             f"{ecarts_attendus[libelle]:+.2f})")
intactes = [libelle for libelle, e in ecarts_attendus.items() if abs(e) <= 0.005]
verifier("les colonnes que rien ne touche n'affichent pas d'écart",
         all(not tuiles_rectif[libelle].delta for libelle in intactes),
         f"({[(l, tuiles_rectif[l].delta) for l in intactes if tuiles_rectif[l].delta]})")

print("\n7 bis. Rectifications — modifier le tableau de droite")
at = lancer("app_pages/rectifications.py", {
    "mois_travail": (2026, 7),
    f"rectif_2026_7": {"edited_rows": {position: {"pay_n": corrige_D.pay_n + 4}},
                       "added_rows": [], "deleted_rows": []},
})
verifier("aucune exception", not at.exception, str(at.exception))
attente = table(at, "Écart au réel", "Rectifié")
verifier("la correction en attente est détaillée", attente is not None)
verifier("elle rappelle la valeur réelle",
         attente["Réel"].iloc[0] == str(reel_D.pay_n), f"({attente['Réel'].iloc[0]})")
verifier("elle rappelle la valeur en vigueur",
         attente["En vigueur"].iloc[0] == str(corrige_D.pay_n))
verifier("elle montre la valeur rectifiée",
         attente["Rectifié"].iloc[0] == str(corrige_D.pay_n + 4))
verifier("l'écart est mesuré par rapport au réel",
         abs(float(attente["Écart au réel"].iloc[0])
             - (corrige_D.pay_n + 4 - reel_D.pay_n)) < 0.01)

recettes = [m for m in at.metric if m.label == "Recettes"]
verifier("le bandeau chiffre l'écart en attente",
         recettes[0].delta and "+" in recettes[0].delta,
         f"({recettes[0].delta})")
verifier("l'écart correspond aux 4 consultations ajoutées",
         recettes[0].delta.replace(" ", "").replace("TND", "")
         == f"{4 * corrige_D.tarifs['t_pay'] + (corrige_D.recettes - reel_D.recettes):+.0f}",
         f"({recettes[0].delta})")

bouton = [b for b in at.button if "Enregistrer" in b.label][0]
verifier("enregistrement bloqué sans motif", bouton.disabled)
at.text_area(key="motif_rectif_2026_7").set_value("recomptage du registre (test)").run()
bouton = [b for b in at.button if "Enregistrer" in b.label][0]
verifier("débloqué avec motif", not bouton.disabled)
bouton.click().run()
verifier("aucune exception à l'enregistrement", not at.exception, str(at.exception))

conn = store.connect()
verifier("une rectification de plus", store.nb_rectifications(conn) == nb_avant + 1)
verifier("le registre corrigé a bougé",
         store.lire_jour_corrige(conn, D).pay_n == corrige_D.pay_n + 4)
verifier("le registre réel n'a pas bougé",
         store.lire_jour(conn, D).pay_n == reel_D.pay_n)
verifier("le motif est attaché",
         store.rectifications_du_jour(conn, D)[0]["motif"]
         == "recomptage du registre (test)")
conn.close()

print("\n7 ter. Rectifications — un mois sans journée clôturée")
at = lancer("app_pages/rectifications.py", {"mois_travail": (2027, 5)})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("le cas est expliqué",
         any("rien à rectifier" in i.value for i in at.info))
verifier("et renvoie vers la saisie",
         any("Mois comptable" in i.value for i in at.info))

print("\n8. Tableau de bord — les deux registres")
at = lancer("app_pages/tableau_bord.py")
verifier("aucune exception", not at.exception, str(at.exception))
verifier("mention des rectifications",
         any("rectification" in c.value for c in at.caption))
corrige_recettes = at.metric[0].value
at.segmented_control(key="registre_bord").select("Réel").run()
verifier("bascule sans exception", not at.exception, str(at.exception))
verifier("le total change entre les deux vues",
         at.metric[0].value != corrige_recettes,
         f"({at.metric[0].value} vs {corrige_recettes})")

print("\n9. Classeur Excel")
at = lancer("app_pages/excel.py")
verifier("aucune exception", not at.exception, str(at.exception))
verifier("journée rectifiée en attente d'écriture",
         any("en attente" in s.value for s in at.subheader))

print("\n10. Tarifs")
at = lancer("app_pages/tarifs.py")
verifier("aucune exception", not at.exception, str(at.exception))

print("\n11. Bordereaux — un mois, son classeur, son historique")
# La page range le bordereau dans le dossier du comptable et lit ce qui s'y
# trouve. Un test ne touche ni à `F:\comptabilité` ni au dossier de travail :
# on déplace les deux racines dans du temporaire.
from compta import bordereau as _bordereau  # noqa: E402
from compta import dossier_comptable as _dc  # noqa: E402
import compta.config as _config  # noqa: E402
from openpyxl import Workbook as _Workbook  # noqa: E402

racine_test = Path(tempfile.mkdtemp())
dossier_test = racine_test / "bordereaux"
_dc.RACINE = racine_test / "comptabilite"
_bordereau.BORDEREAUX_DIR = dossier_test
_bordereau.ARCHIVES_DIR = racine_test / "archives"
_config.BORDEREAUX_DIR = dossier_test


def transmis_bidon(chemin, annee, mois, jours, ca_jour, patients_jour):
    """Un classeur mensuel du comptable, en-têtes sur trois lignes et valeurs
    en dur — comme ceux qu'Excel enregistre."""
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb = _Workbook()
    ws = wb.active
    ws["B1"], ws["H1"], ws["H2"] = "DATE", "CA A ", "DECLARER"
    ws["G2"], ws["G3"] = "patients à ", "DECLARER"
    for n in range(jours):
        r = 4 + n
        ws.cell(row=r, column=2, value=dt.datetime(annee, mois, n + 1))
        ws.cell(row=r, column=7, value=patients_jour)
        ws.cell(row=r, column=8, value=ca_jour)
    wb.save(chemin)
    wb.close()

# un mois qui porte des chiffres, pour que le bouton ne soit pas désactivé
conn = store.connect()
mois_plein = conn.execute(
    "SELECT substr(date,1,7) m FROM jours GROUP BY m HAVING SUM(recettes) > 0 "
    "ORDER BY m DESC LIMIT 1"
).fetchone()["m"]
conn.close()
annee_test, mois_test = int(mois_plein[:4]), int(mois_plein[5:7])

# sans mois choisi, la page s'ouvre sur le dernier mois qui porte des chiffres —
# pas sur le mois à saisir, qui n'offrirait qu'un bouton grisé. La page
# d'accueil a déjà résolu le mois en passant : on l'efface pour arriver ici sans
# choix, sinon le test ne mesure que l'accord fortuit des deux défauts.
at = lancer()
at.session_state["mois_travail"] = None
at.switch_page("app_pages/bordereaux.py").run()
conn = store.connect()
dernier_rempli = store.dernier_jour_rempli(conn)
conn.close()
verifier("sans mois choisi, elle s'ouvre sur le dernier mois renseigné",
         at.session_state["mois_travail"]
         == (dernier_rempli.year, dernier_rempli.month),
         f"({at.session_state['mois_travail']} au lieu de "
         f"{(dernier_rempli.year, dernier_rempli.month)})")
verifier("et le bouton y est donc actif",
         (b := next((b for b in at.button if "Générer" in b.label), None))
         is not None and not b.disabled)

at = lancer("app_pages/bordereaux.py",
            {"mois_travail": (annee_test, mois_test)})
verifier("un mois choisi ailleurs est respecté",
         at.session_state["mois_travail"] == (annee_test, mois_test),
         f"({at.session_state['mois_travail']})")
verifier("aucune exception", not at.exception, str(at.exception))

# Arriver sur un mois vide ne doit pas être un cul-de-sac : c'est ce que fait
# l'accueil, qui ouvre sur le mois **à saisir**. On prend un mois franchement
# vide plutôt que `mois_a_saisir`, que les sections précédentes ont pu remplir.
conn = store.connect()
dernier_rempli = store.dernier_jour_rempli(conn)
mois_vide = (dernier_rempli.year + 1, dernier_rempli.month)
reste_vide = conn.execute(
    "SELECT COUNT(*) n FROM jours WHERE substr(date,1,7)=? AND recettes>0",
    (f"{mois_vide[0]:04d}-{mois_vide[1]:02d}",),
).fetchone()["n"]
conn.close()
verifier("le mois témoin est bien sans chiffres", reste_vide == 0,
         f"({reste_vide} journée(s) renseignée(s))")

at_vide = lancer("app_pages/bordereaux.py", {"mois_travail": mois_vide})
verifier("un mois vide le dit",
         any("aucun chiffre" in w.value for w in at_vide.warning),
         str(at_vide.warning))
verifier("le bouton de génération y est grisé",
         (b := next((b for b in at_vide.button if "Générer" in b.label), None))
         is not None and b.disabled)
retour = next((b for b in at_vide.button
               if "dernier mois renseigné" in b.label), None)
verifier("et le dernier mois renseigné est à un clic", retour is not None,
         f"({[b.label for b in at_vide.button]})")
if retour is not None:
    at_vide = retour.click().run()
    verifier("le clic y mène",
             at_vide.session_state["mois_travail"]
             == (dernier_rempli.year, dernier_rempli.month),
             f"({at_vide.session_state['mois_travail']})")
    verifier("et le bouton de génération y redevient actif",
             (b := next((b for b in at_vide.button if "Générer" in b.label),
                        None)) is not None and not b.disabled)
verifier("la page annonce les bordereaux",
         any("Bordereaux" in h.value for h in at.header))
verifier("le mois affiché est celui demandé",
         at.session_state["mois_travail"] == (annee_test, mois_test),
         f"({at.session_state['mois_travail']})")
verifier("les recettes du mois sont affichées", metrique(at, "Recettes") is not None)
verifier("le CA à déclarer aussi", metrique(at, "CA à déclarer") is not None)
for poste in ("Caisse", "À la charge CNAM", "À verser à la banque", "Patients"):
    verifier(f"« {poste} » est au bandeau", metrique(at, poste) is not None)
detail = table(at, "Recettes", "À verser banque")
verifier("le détail du mois montre le à-verser-à-la-banque",
         detail is not None and len(detail) > 0,
         "(colonne absente du tableau des journées)")
# l'aperçu doit porter les 28 colonnes du fichier, pas une sélection
libelles = [libelle for _, libelle, _ in _bordereau.APERCU]
manquantes = [c for c in libelles
              if detail is None or c not in detail.columns]
verifier("aucune colonne du bordereau ne manque à l'aperçu",
         not manquantes, f"({manquantes})")
verifier("les totaux hors bandeau sont dits quelque part",
         any("TVA déduite" in c.value and "dépenses" in c.value
             for c in at.caption),
         "(TVA déduite, non déclaré, laser, divers, dépenses)")
verifier("aucun bordereau n'est encore édité",
         any("pas encore été édité" in c.value for c in at.caption),
         "(historique non vide au départ)")

bouton = next((b for b in at.button if "Générer" in b.label), None)
verifier("le bouton de génération est là", bouton is not None)
verifier("il est actif sur un mois renseigné",
         bouton is not None and not bouton.disabled)

at = bouton.click().run()
verifier("générer ne lève rien", not at.exception, str(at.exception))
verifier("le fichier est annoncé enregistré",
         any("enregistré" in s.value for s in at.success), str(at.success))
attendu_dossier = (_dc.RACINE / str(annee_test)
                   / _dc.nom_dossier(annee_test, mois_test))
produit = attendu_dossier / _bordereau.nom_fichier(annee_test, mois_test)
verifier("le bordereau est rangé dans le dossier du mois, chez le comptable",
         produit.exists() and produit.stat().st_size > 0, f"({produit})")
verifier("le dossier du mois a été créé au besoin", attendu_dossier.is_dir())
verifier("il se télécharge depuis la page",
         any("Télécharger" in d.label for d in at.download_button))

historique = table(at, "Édité le", "Recettes", "Δ recettes")
verifier("l'édition figure à l'historique",
         historique is not None and len(historique) == 1,
         f"({None if historique is None else len(historique)} ligne(s))")
verifier("l'historique dit de quel registre elle vient",
         historique is not None and historique["Registre"].iloc[0] == "corrigé",
         f"({None if historique is None else historique['Registre'].iloc[0]})")

# rééditer : une ligne de plus, la première intacte
at = lancer("app_pages/bordereaux.py", {"mois_travail": (annee_test, mois_test)})
bouton = next((b for b in at.button if "Générer" in b.label), None)
at = bouton.click().run()
verifier("rééditer ne lève rien", not at.exception, str(at.exception))
historique = table(at, "Édité le", "Recettes", "Δ recettes")
verifier("la réédition s'ajoute à l'historique",
         historique is not None and len(historique) == 2,
         f"({None if historique is None else len(historique)} ligne(s))")
verifier("l'édition précédente est archivée dans le projet, pas chez le comptable",
         _bordereau.ARCHIVES_DIR.exists()
         and any(_bordereau.ARCHIVES_DIR.glob("*.xlsx")))
verifier("le dossier du comptable ne garde qu'un bordereau",
         len(list(attendu_dossier.glob("bordereau-*.xlsx"))) == 1,
         f"({[p.name for p in attendu_dossier.glob('bordereau-*.xlsx')]})")
verifier("la page signale l'édition précédente",
         any("déjà été édité" in i.value for i in at.info), str(at.info))

print("\n11 bis. Bordereaux — ce qui a été transmis au comptable")
# un fichier transmis pour le mois de test, avec un CA volontairement plus bas
transmis_bidon(attendu_dossier / f"{mois_test:02d}-{annee_test} transmis.xlsx",
               annee_test, mois_test, 28, 100.0, 3)
at = lancer("app_pages/bordereaux.py", {"mois_travail": (annee_test, mois_test)})
verifier("aucune exception", not at.exception, str(at.exception))
verifier("la page dit que ce ne sont pas le registre",
         any("pas le registre" in c.value for c in at.caption), "(avertissement absent)")
# le libellé dit « Lire » ou « Relire » selon qu'un passage a déjà eu lieu — et
# la base de test est une copie de la vraie, qui peut en porter un
lire_dossier = next((b for b in at.button if "dossier" in b.label), None)
verifier("le dossier du comptable se lit depuis la page",
         lire_dossier is not None, f"({[b.label for b in at.button]})")

at = lire_dossier.click().run()
verifier("la lecture ne lève rien", not at.exception, str(at.exception))
verifier("elle rend compte de ce qu'elle a vu",
         any("classeurs vus" in s.value for s in at.success), str(at.success))

conn = store.connect()
par_mois = store.transmis_par_mois(conn)
conn.close()
cle_test = f"{annee_test}-{mois_test:02d}"
verifier("le mois transmis est au cache", cle_test in par_mois,
         f"({sorted(par_mois)})")
verifier("son CA est celui du fichier, pas celui du registre",
         cle_test in par_mois
         and abs(par_mois[cle_test]["totaux"]["ca_declare"] - 2800.0) < 0.5,
         f"({par_mois.get(cle_test, {}).get('totaux')})")
verifier("notre propre bordereau n'a pas été pris pour un fichier transmis",
         all("bordereau-" not in f["chemin"] for f in par_mois.values()),
         f"({[f['chemin'] for f in par_mois.values()]})")
verifier("la comparaison s'affiche",
         metrique(at, "CA à déclarer transmis") is not None,
         f"({[m.label for m in at.metric]})")
verifier("elle montre le chiffre du fichier",
         (m := metrique(at, "CA à déclarer transmis")) is not None
         and "2 800" in m.value, f"({m.value if m else None})")
verifier("l'écart est lisible : milliers espacés et unité",
         m is not None and m.delta.endswith(" TND") and "," not in m.delta,
         f"({m.delta if m else None})")
verifier("un graphe d'évolution est tracé", len(at.get("vega_lite_chart")) > 0)
# le choix des mois tracés se vérifie sur `bordereau.mois_a_tracer`,
# dans test_bordereau : ici on s'assure seulement que le graphe existe

print(f"\n{'=' * 50}\n{ok} réussis, {ko} échecs\n{'=' * 50}")
sys.exit(1 if ko else 0)
