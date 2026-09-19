"""Bordereau mensuel : un mois du registre, extrait dans un classeur autonome.

Le bordereau reprend la feuille de « compta jdid .xlsx » telle quelle — colonnes
A→AD, en-tête sur trois lignes, données à partir de la ligne 4 — mais sur un seul
mois et dans un fichier neuf.

Le point qui fait tout : **les montants restent des formules**. Seuls les
compteurs et les montants saisis sont écrits en valeur (C, E, H, J, L, N, P puis
R, S, T, V) ; D, F, G, I, K, M, O, Q, U, W, X, Y, Z, AA, AB, AC, AD sont les
formules du classeur, réécrites au numéro de ligne du bordereau. Corriger un
compteur dans le bordereau met donc ses totaux à jour tout seul, exactement comme
dans le classeur. La ligne de total en bas est un `SUM` sur chaque colonne.

Chaque journée porte **son** tarif, celui que le registre lui connaît : une
journée de 2019 s'écrit au prix de 2019. C'est la même règle que partout ailleurs
dans l'app, et elle est ici visible à l'œil nu dans la formule (`=C12*52`).

Le classeur du cabinet n'est jamais ouvert : on ne fait que lire le registre.

Une seule liberté prise sur le modèle : la colonne **Y — « à la charge STAR »**.
Le classeur l'annonce en en-tête mais ne la remplit plus depuis 2016 ; le
bordereau la remplit (`=K{ligne}`), une colonne titrée et vide n'étant lisible
par personne. Aucun total ne dépend de Y, ni dans le classeur ni ici.
"""

from __future__ import annotations

import datetime as dt
import hashlib
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import quote_sheetname
from openpyxl.worksheet.properties import PageSetupProperties

from .config import (ARCHIVES_DIR, BORDEREAUX_DIR, JOURS_FR, MOIS_FR,
                     SAISIE_COLS)
from .model import Journee, LIBELLES_TARIFS, mois_long

# le classeur réserve trois lignes d'en-tête ; les journées commencent en 4
PREMIERE_LIGNE = 4

# En-têtes repris mot pour mot du classeur, espaces parasites en moins : le
# comptable doit retrouver la feuille qu'il reçoit d'habitude.
ENTETES: dict[str, str] = {
    "B1": "DATE", "C1": "C° CNAM", "J1": "assurances", "L1": "payant",
    "N1": "cataractes", "R1": "laser", "U1": "RECETTES", "V1": "DEPENSES",
    "Y1": "à la charge", "Z1": "a la charge", "AA1": "a verser",
    "AC1": "CA A", "AD1": "TVA DED",

    "C2": "APCI", "E2": "TP", "H2": "remb", "K2": "total star",
    "L2": "NBR", "M2": "MONT", "N2": "CNAM", "O2": "total",
    "P2": "payant", "Q2": "total", "S2": "non", "T2": "divers RECU",
    "W2": "TVA TP", "X2": "caisse", "AA2": "a", "AB2": "patients à",
    "AC2": "DECLARER",

    "C3": "NBR", "D3": "MONT", "E3": "NBR", "F3": "recu", "G3": "reste",
    "H3": "NBR", "I3": "MONT", "R3": "declaré", "S3": "declaré",
    "Y3": "STAR", "Z3": "CNAM", "AA3": "la banque", "AB3": "DECLARER",
}

FUSIONS = [
    "A1:A3", "B1:B3", "C1:I1", "C2:D2", "E2:G2", "H2:I2", "J1:K1", "K2:K3",
    "L1:M1", "L2:L3", "M2:M3", "N1:Q1", "R1:S1", "T2:T3", "U1:U3", "V1:V3",
    "W2:W3", "X2:X3", "Y1:Y2", "Z1:Z2", "AD1:AD3",
]

# toutes les colonnes chiffrées, dans l'ordre : ce sont elles que la ligne de
# total additionne
COLONNES = [
    "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q",
    "R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AD",
]

# compteurs d'un côté, montants de l'autre : ils ne s'affichent pas pareil
COMPTEURS_COL = {"C", "E", "H", "J", "L", "N", "P", "AB"}

LARGEURS = {"A": 11, "B": 11.5, "AD": 10}

BLEU = "2A78D6"
ARDOISE = "44515E"
GRIS = "F2F4F6"
TRAIT = "D4DAE0"


def _fmt(t: float) -> str:
    """Un tarif tel qu'il s'écrit dans une formule : 72, pas 72.0."""
    return str(int(t)) if float(t).is_integer() else str(t)


def nom_feuille(annee: int, mois: int) -> str:
    return mois_long(annee, mois)


def nom_fichier(annee: int, mois: int) -> str:
    """`bordereau-2026-08-aout.xlsx` : se trie tout seul, se lit quand même."""
    sans_accent = (MOIS_FR[mois - 1]
                   .replace("û", "u").replace("é", "e").replace("è", "e"))
    return f"bordereau-{annee}-{mois:02d}-{sans_accent}.xlsx"


# Les 28 colonnes chiffrées, lisibles : lettre du classeur, libellé, et
# l'attribut de `Journee` qui porte la valeur. C'est ce qui permet d'afficher le
# bordereau à l'écran **sans en perdre une colonne** — un aperçu qui oublierait
# les dépenses ou le laser montrerait autre chose que ce qui part.
# `COLONNES` et cette liste se tiennent : un test le vérifie.
APERCU: list[tuple[str, str, str]] = [
    ("C", "APCI", "apci_n"),
    ("D", "APCI mont", "apci_m"),
    ("E", "TP", "tp_n"),
    ("F", "TP reçu", "tp_recu"),
    ("G", "TP reste", "tp_reste"),
    ("H", "Remb", "remb_n"),
    ("I", "Remb mont", "remb_m"),
    ("J", "Assurance", "star_n"),
    ("K", "Assurance mont", "star_m"),
    ("L", "Payant", "pay_n"),
    ("M", "Payant mont", "pay_m"),
    ("N", "Cat. CNAM", "cat_cnam_n"),
    ("O", "Cat. CNAM mont", "cat_cnam_m"),
    ("P", "Cat. payante", "cat_pay_n"),
    ("Q", "Cat. payante mont", "cat_pay_m"),
    ("R", "Laser déclaré", "laser_dec"),
    ("S", "Laser non déclaré", "laser_nd"),
    ("T", "Divers", "divers"),
    ("U", "Recettes", "recettes"),
    ("V", "Dépenses", "depenses"),
    ("W", "TVA TP", "tva_tp"),
    ("X", "Caisse", "caisse"),
    ("Y", "À la charge STAR", "star_m"),
    ("Z", "À la charge CNAM", "charge_cnam"),
    ("AA", "À verser banque", "banque"),
    ("AB", "Patients", "patients"),
    ("AC", "CA à déclarer", "ca_declare"),
    ("AD", "TVA déduite", "tva_deduite"),
]

# les colonnes qui comptent des actes ou des patients, pas des dinars
APERCU_COMPTEURS = {libelle for lettre, libelle, _ in APERCU
                    if lettre in COMPTEURS_COL}


# --- totaux ------------------------------------------------------------------

POSTES_TOTAUX: list[tuple[str, str, str]] = [
    # clé interne, libellé, colonne du classeur
    ("recettes", "Recettes", "U"),
    ("ca_declare", "CA à déclarer", "AC"),
    ("tva_deduite", "TVA déduite", "AD"),
    ("caisse", "Caisse", "X"),
    ("charge_cnam", "À la charge CNAM", "Z"),
    ("banque", "À verser à la banque", "AA"),
    ("laser_dec", "Laser déclaré", "R"),
    ("laser_nd", "Laser non déclaré", "S"),
    ("divers", "Divers reçu", "T"),
    ("depenses", "Dépenses", "V"),
    ("patients", "Patients", "AB"),
    ("cat_cnam_n", "Cataractes CNAM", "N"),
    ("cat_pay_n", "Cataractes payantes", "P"),
]


def totaux(journees: list[Journee]) -> dict[str, float]:
    """Les totaux du mois, calculés par le miroir Python des formules.

    Ce sont les mêmes nombres que ceux qu'Excel affichera en ouvrant le
    bordereau : `model.py` reproduit les formules du classeur.
    """
    somme = lambda champ: float(sum(getattr(j, champ) for j in journees))  # noqa: E731
    resultat = {cle: somme(cle) for cle, _, _ in POSTES_TOTAUX}
    resultat |= {
        "non_declare": somme("recettes") - somme("ca_declare"),
        "jours": float(len(journees)),
        "jours_actifs": float(sum(1 for j in journees if j.recettes > 0)),
        "actes_chirurgie": somme("actes_chirurgie"),
    }
    return resultat


def tarifs_du_mois(journees: list[Journee]) -> dict[str, list[float]]:
    """Les tarifs rencontrés dans le mois, poste par poste.

    Plusieurs valeurs pour un poste : le tarif a changé en cours de mois, ou le
    mois mélange des journées d'époques différentes. Le bordereau le montre
    plutôt que d'en choisir un.
    """
    releve: dict[str, list[float]] = {}
    for j in journees:
        for cle, valeur in j.tarifs.items():
            vus = releve.setdefault(cle, [])
            if valeur not in vus:
                vus.append(valeur)
    return {cle: sorted(vus) for cle, vus in releve.items()}


# --- fabrication du classeur -------------------------------------------------

def _entetes(ws) -> None:
    fond = PatternFill("solid", fgColor=BLEU)
    police = Font(bold=True, color="FFFFFF", size=9)
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in (1, 2, 3):
        for col in ["A", "B"] + COLONNES:
            cellule = ws[f"{col}{r}"]
            cellule.fill = fond
            cellule.font = police
            cellule.alignment = centre
        ws.row_dimensions[r].height = 16

    for coordonnee, texte in ENTETES.items():
        ws[coordonnee] = texte
    for plage in FUSIONS:
        ws.merge_cells(plage)


def _journee(ws, r: int, j: Journee) -> None:
    """Une ligne de journée : les saisies en valeur, le reste en formules."""
    t = j.tarifs
    ws[f"A{r}"] = JOURS_FR[j.date.weekday()]
    date = ws[f"B{r}"]
    date.value = dt.datetime(j.date.year, j.date.month, j.date.day)
    date.number_format = "DD/MM/YYYY"

    for cle, col in SAISIE_COLS.items():
        ws[f"{col}{r}"] = getattr(j, cle)

    ws[f"D{r}"] = f"=C{r}*{_fmt(t['t_apci'])}"
    ws[f"F{r}"] = f"=E{r}*{_fmt(t['t_tp_recu'])}"
    ws[f"G{r}"] = f"=E{r}*{_fmt(t['t_tp_reste'])}"
    ws[f"I{r}"] = f"=H{r}*{_fmt(t['t_remb'])}"
    ws[f"K{r}"] = f"=J{r}*{_fmt(t['t_star'])}"
    ws[f"M{r}"] = f"=L{r}*{_fmt(t['t_pay'])}"
    ws[f"O{r}"] = f"=N{r}*{_fmt(t['t_cat_cnam'])}"
    ws[f"Q{r}"] = f"=P{r}*{_fmt(t['t_cat_pay'])}"
    ws[f"W{r}"] = f"=E{r}*{_fmt(t['t_tva_tp'])}"
    ws[f"U{r}"] = f"=D{r}+F{r}+I{r}+K{r}+M{r}+O{r}+Q{r}+S{r}+T{r}+R{r}+W{r}"
    ws[f"X{r}"] = f"=SUM(F{r},I{r},M{r},T{r},W{r})-V{r}"
    ws[f"Y{r}"] = f"=K{r}"
    ws[f"Z{r}"] = f"=SUM(D{r},G{r},O{r})"
    ws[f"AA{r}"] = f"=SUM(F{r},I{r},K{r},M{r},O{r}/2)"
    ws[f"AB{r}"] = f"=C{r}+E{r}+H{r}+J{r}+L{r}"
    ws[f"AC{r}"] = f"=D{r}+F{r}+G{r}+I{r}+K{r}+M{r}+O{r}+R{r}"
    ws[f"AD{r}"] = f"=AC{r}-(AC{r}/1.06)"


def _mise_en_forme(ws, premiere: int, derniere: int) -> None:
    bordure = Border(bottom=Side(style="thin", color=TRAIT))
    for r in range(premiere, derniere + 1):
        dimanche = ws[f"A{r}"].value == "DIMANCHE"
        for col in ["A", "B"] + COLONNES:
            cellule = ws[f"{col}{r}"]
            cellule.border = bordure
            if col in COMPTEURS_COL:
                cellule.number_format = "0"
            elif col == "AD":
                cellule.number_format = "# ##0.00"
            elif col not in ("A", "B"):
                cellule.number_format = "# ##0"
            # le dimanche se repère d'un coup d'œil : c'est la ligne à zéro
            # qui n'est pas un oubli
            if dimanche:
                cellule.fill = PatternFill("solid", fgColor=GRIS)
        ws[f"A{r}"].font = Font(size=9, color="666666" if dimanche else "000000")


def _total(ws, r: int, premiere: int, derniere: int, periode: str) -> None:
    ws[f"A{r}"] = "TOTAL"
    ws[f"B{r}"] = periode
    haut = Border(top=Side(style="medium", color=ARDOISE))
    for col in COLONNES:
        ws[f"{col}{r}"] = f"=SUM({col}{premiere}:{col}{derniere})"
    for col in ["A", "B"] + COLONNES:
        cellule = ws[f"{col}{r}"]
        cellule.font = Font(bold=True, size=9)
        cellule.border = haut
        cellule.fill = PatternFill("solid", fgColor=GRIS)
        if col in COMPTEURS_COL:
            cellule.number_format = "0"
        elif col == "AD":
            cellule.number_format = "# ##0.00"
        elif col not in ("A", "B"):
            cellule.number_format = "# ##0"


def _recapitulatif(wb, feuille: str, ligne_total: int, annee: int, mois: int,
                   journees: list[Journee], registre: str,
                   edite_le: dt.datetime) -> None:
    """Deuxième feuille : d'où vient ce bordereau, et ce qu'il totalise.

    Les totaux y sont des **renvois** vers la ligne de total du mois, pas des
    nombres recopiés : corriger une journée met le récapitulatif à jour.
    """
    ws = wb.create_sheet("Récapitulatif")
    reference = quote_sheetname(feuille)

    ws["A1"] = f"Bordereau mensuel — {mois_long(annee, mois)}"
    ws["A1"].font = Font(bold=True, size=14, color=ARDOISE)
    ws["A2"] = "Cabinet d'ophtalmologie"
    ws["A2"].font = Font(italic=True, color="666666")

    debut, fin = min(j.date for j in journees), max(j.date for j in journees)
    provenance = [
        ("Période", f"du {debut:%d/%m/%Y} au {fin:%d/%m/%Y} "
                    f"({len(journees)} journées)"),
        ("Registre", f"registre {registre}"),
        ("Source", "compta jdid .xlsx — registre du cabinet"),
        ("Édité le", edite_le.strftime("%d/%m/%Y à %H:%M")),
    ]
    ligne = 4
    for libelle, valeur in provenance:
        ws[f"A{ligne}"] = libelle
        ws[f"A{ligne}"].font = Font(bold=True, size=9)
        ws[f"B{ligne}"] = valeur
        ligne += 1

    ligne += 1
    ws[f"A{ligne}"] = "Totaux du mois"
    ws[f"A{ligne}"].font = Font(bold=True, size=11, color=ARDOISE)
    ligne += 1
    for _, libelle, colonne in POSTES_TOTAUX:
        ws[f"A{ligne}"] = libelle
        cellule = ws[f"B{ligne}"]
        cellule.value = f"={reference}!{colonne}{ligne_total}"
        cellule.number_format = "0" if colonne in COMPTEURS_COL else "# ##0"
        ws[f"C{ligne}"] = f"colonne {colonne}"
        ws[f"C{ligne}"].font = Font(size=8, color="999999")
        ligne += 1

    ligne += 1
    ws[f"A{ligne}"] = "Tarifs appliqués"
    ws[f"A{ligne}"].font = Font(bold=True, size=11, color=ARDOISE)
    ligne += 1
    for cle, valeurs in tarifs_du_mois(journees).items():
        ws[f"A{ligne}"] = LIBELLES_TARIFS.get(cle, cle)
        ws[f"B{ligne}"] = " · ".join(_fmt(v) for v in valeurs)
        if len(valeurs) > 1:
            ws[f"C{ligne}"] = "tarif non uniforme sur le mois"
            ws[f"C{ligne}"].font = Font(size=8, color="B35A00")
        ligne += 1

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 28


def construire(journees: list[Journee], annee: int, mois: int,
               registre: str = "corrigé",
               edite_le: dt.datetime | None = None) -> bytes:
    """Le bordereau du mois, en octets, prêt à enregistrer ou à télécharger.

    `journees` est le mois entier, dans l'ordre, telles que le registre les
    donne — chacune avec ses tarifs.
    """
    if not journees:
        raise ValueError("Un bordereau sans journée ne veut rien dire.")

    edite_le = edite_le or dt.datetime.now()
    feuille = nom_feuille(annee, mois)

    wb = Workbook()
    ws = wb.active
    ws.title = feuille

    _entetes(ws)

    r = PREMIERE_LIGNE
    for j in sorted(journees, key=lambda x: x.date):
        _journee(ws, r, j)
        r += 1
    derniere = r - 1

    _mise_en_forme(ws, PREMIERE_LIGNE, derniere)
    ligne_total = derniere + 1
    _total(ws, ligne_total, PREMIERE_LIGNE, derniere, mois_long(annee, mois))

    for col in ["A", "B"] + COLONNES:
        ws.column_dimensions[col].width = LARGEURS.get(col, 8.5)
    ws.freeze_panes = ws[f"C{PREMIERE_LIGNE}"]
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    # 28 colonnes sur du A4 : sans cette propriété, Excel imprime la moitié du
    # bordereau sur une deuxième page
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    _recapitulatif(wb, feuille, ligne_total, annee, mois, journees, registre,
                   edite_le)

    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def empreinte(contenu: bytes) -> str:
    """Signature courte d'un bordereau : deux éditions identiques se voient."""
    return hashlib.sha256(contenu).hexdigest()[:16]


# --- enregistrement sur le disque --------------------------------------------

def destination(annee: int, mois: int, racine: Path | None = None,
                creer: bool = False) -> Path:
    """Où va le bordereau d'un mois : le dossier du comptable, s'il existe.

    Le cabinet range déjà ses mois dans `F:\\comptabilité\\<année>\\<mois>\\` — le
    bordereau y rejoint le reste du mois plutôt que de vivre à part. À défaut de
    dossier d'année (un mois à venir, un disque absent), on retombe sur
    `bordereaux/<année>/` dans le projet, et rien n'est perdu.
    """
    from . import dossier_comptable

    trouve = dossier_comptable.dossier_du_mois(annee, mois, racine, creer=creer)
    if trouve is not None:
        return trouve
    return (Path(racine) if racine else BORDEREAUX_DIR) / str(annee)


def enregistrer(contenu: bytes, annee: int, mois: int,
                dossier: Path | None = None,
                racine_comptable: Path | None = None) -> tuple[Path, Path | None]:
    """Pose le bordereau du mois à son adresse, l'édition précédente à l'abri.

    Sans `dossier`, il va dans le dossier du mois chez le comptable — créé s'il
    le faut. Son nom, `bordereau-2026-08-aout.xlsx`, ne ressemble à aucun de
    ceux qui s'y trouvent : **il ne recouvre jamais un fichier existant**.

    Un mois a toujours **un** bordereau courant, au même chemin — c'est lui
    qu'on envoie. Rééditer ne l'écrase pas en silence : l'ancien part dans
    `bordereaux/archives/` du projet, horodaté, comme les sauvegardes du
    classeur. Les archives restent chez nous : le dossier du comptable garde
    le seul bordereau en cours.

    Retourne le chemin écrit et celui de l'archive, s'il y en a eu une.
    """
    if dossier is not None:
        cible = Path(dossier) / str(annee)
    else:
        cible = destination(annee, mois, racine_comptable, creer=True)
    cible.mkdir(parents=True, exist_ok=True)
    chemin = cible / nom_fichier(annee, mois)

    archive = None
    if chemin.exists():
        archives = (Path(dossier) / "archives" if dossier is not None
                    else ARCHIVES_DIR)
        archives.mkdir(parents=True, exist_ok=True)
        horodatage = dt.datetime.fromtimestamp(chemin.stat().st_mtime)
        archive = archives / f"{chemin.stem}_{horodatage:%Y%m%d_%H%M%S}.xlsx"
        # deux éditions dans la même seconde : on ne perd pas la première
        suffixe = 1
        while archive.exists():
            suffixe += 1
            archive = archives / f"{chemin.stem}_{horodatage:%Y%m%d_%H%M%S}-{suffixe}.xlsx"
        chemin.replace(archive)

    chemin.write_bytes(contenu)
    return chemin, archive
